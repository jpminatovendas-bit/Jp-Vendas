"""JP Minato: aplicativo local de contatos e catálogos municipais."""
from __future__ import annotations
import contextvars
import operators131
import galleries135
import inbox131
import dashboard132
import communication
import send_activity130
import whatsapp_sender129
import smtp_mail127
import whatsapp_send126
import catalogue133 as catalogue125
import pdf_output128
import editing_activity
import time, colorsys, math, io, json, mimetypes, os, re, secrets, shutil, sqlite3, sys, subprocess, threading, unicodedata, webbrowser, zipfile
from datetime import datetime
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote
import fitz
import openpyxl
from tse import read_candidates,clean
import contact_extractor as contact_extractor
import municipal_media
import whatsapp_verify
import restoration_ai
import crest_capture
import capture_routes
import google_contacts
import whatsapp_browser
import remote_access
import data_storage
from whatsapp_links import parse_link

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'dados'
IMAGE_EXT = {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}
UF_DDD = {'AC':['68'],'AL':['82'],'AP':['96'],'AM':['92','97'],'BA':['71','73','74','75','77'],'CE':['85','88'],'DF':['61'],'ES':['27','28'],'GO':['62','64'],'MA':['98','99'],'MT':['65','66'],'MS':['67'],'MG':['31','32','33','34','35','37','38'],'PA':['91','93','94'],'PB':['83'],'PR':['41','42','43','44','45','46'],'PE':['81','87'],'PI':['86','89'],'RJ':['21','22','24'],'RN':['84'],'RS':['51','53','54','55'],'RO':['69'],'RR':['95'],'SC':['47','48','49'],'SP':['11','12','13','14','15','16','17','18','19'],'SE':['79'],'TO':['63']}
EXCEPTIONS = {('SC','PORTOUNIAO'): ['42']}

def norm(value):
    return re.sub(r'[^A-Z0-9]', '', unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().upper())

def now():
    return datetime.now().isoformat(timespec='seconds')

def useful(value):
    s = str(value or '').strip()
    return bool(s) and not any(x in norm(s) for x in ['NAODISPONIVEL','NAOIDENTIFICADO','NAOENCONTRADO','NAOINFORMADO','SEMCONTATO'])

def ddds(value):
    found = re.findall(r'\((\d{2})\)\s*\d', str(value or ''))
    if found:
        return list(dict.fromkeys(found))
    result = []
    for piece in re.split(r'[,;|\n]', str(value or '')):
        digits = re.sub(r'\D', '', piece)
        if len(digits) in (12, 13) and digits.startswith('55'):
            digits = digits[2:]
        if len(digits) in (10, 11):
            result.append(digits[:2])
    return list(dict.fromkeys(result))

def safe_filename(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '', str(value)).strip(' .')[:120] or 'SEM_NOME'

def normalize_phone(value):
    text=str(value or '').strip()
    if not re.fullmatch(r'[+\d\s().-]+',text):raise ValueError('Cole apenas um número de telefone com DDD.')
    number=re.sub(r'\D','',text)
    if len(number) in (12,13) and number.startswith('55'):number=number[2:]
    if len(number) not in (10,11):raise ValueError('Informe DDD e telefone, com 10 ou 11 dígitos, ou prefixo +55.')
    if number[:2] not in {d for values in UF_DDD.values() for d in values}:raise ValueError('DDD brasileiro inválido.')
    subscriber=number[2:]
    if not (len(subscriber)==8 and subscriber[0] in '2345' or len(subscriber)==9 and subscriber[0]=='9'):
        raise ValueError('Confira o número: fixo tem 8 dígitos iniciados em 2–5; celular tem 9, iniciado em 9.')
    return '55'+number

def phone_numbers(value):
    found=[]
    for match in re.finditer(r'(?<!\d)(?:\+?55[\s.-]*)?\(?\d{2}\)?[\s.-]*\d{4,5}[\s.-]*\d{4}(?!\d)',str(value or '')):
        try:found.append(normalize_phone(match.group()))
        except ValueError:pass
    return list(dict.fromkeys(found))

class Store:
    def __init__(self, folder=DATA):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.db_path = self.folder / 'empresa.sqlite'
        self.lock = threading.RLock()
        self.files = {}
        self.jobs = {}
        self.capture_tickets = {}
        self.capture_sessions = {}
        self.capture_sequence = 0
        self.capture_events = []
        self.capture_updates = []
        self.capture_condition = threading.Condition(self.lock)
        self.job_lock = threading.Lock()
        self.import_lock = threading.Lock()
        self.google_sync_lock = threading.Lock()
        with self.connect() as db:
            db.execute('PRAGMA journal_mode=WAL').fetchone()
            db.execute('CREATE TABLE IF NOT EXISTS commercial_routes (id TEXT PRIMARY KEY, value TEXT NOT NULL)')
            db.executescript('CREATE TABLE IF NOT EXISTS municipalities (key TEXT PRIMARY KEY, source TEXT NOT NULL, overrides TEXT NOT NULL DEFAULT "{}", local TEXT NOT NULL DEFAULT "{}"); CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT); CREATE TABLE IF NOT EXISTS projects (key TEXT PRIMARY KEY, value TEXT); CREATE TABLE IF NOT EXISTS interactions (id TEXT PRIMARY KEY, municipality_key TEXT NOT NULL, timestamp TEXT NOT NULL, kind TEXT NOT NULL, channel TEXT NOT NULL, target TEXT NOT NULL, recipient TEXT NOT NULL, pdf_path TEXT, pdf_name TEXT, notes TEXT NOT NULL DEFAULT ""); CREATE INDEX IF NOT EXISTS interactions_city ON interactions(municipality_key,timestamp);')

        with self.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS interaction_corrections (id TEXT PRIMARY KEY, active INTEGER NOT NULL, audit TEXT NOT NULL)')
            db.execute('CREATE TABLE IF NOT EXISTS interaction_files (id TEXT PRIMARY KEY, sha256 TEXT NOT NULL)')
            whatsapp_send126.initialize(db)
        whatsapp_sender129.initialize(self)
        smtp_mail127.initialize(self)
        operators131.initialize(self)
        inbox131.initialize(self)
        dashboard132.initialize(self)
        municipal_media.initialize(self)
        contact_extractor.initialize(self)
        whatsapp_verify.initialize(self)
        whatsapp_browser.initialize(self)
        restoration_ai.initialize(self)

    def route_identity(self,db):
        records=[(k,json.loads(v)) for k,v in db.execute('SELECT id,value FROM commercial_routes ORDER BY rowid')]
        counter=db.execute("SELECT value FROM settings WHERE key='route_counter'").fetchone()
        next_number=max([int(json.loads(counter[0])) if counter else 0]+[(r.get('number') or 0) for _,r in records])
        colors={r.get('color') for _,r in records if r.get('color')}
        for key,r in records:
            if not r.get('number'):
                next_number+=1;r['number']=next_number
            if not r.get('color'):
                hue=((r['number']-1)*137.508)%360
                while True:
                    rgb=colorsys.hls_to_rgb(hue/360,.38,.70);color='#'+''.join(f'{round(v*255):02x}' for v in rgb)
                    if color not in colors:break
                    hue=(hue+.731)%360
                r['color']=color;colors.add(color)
            db.execute('UPDATE commercial_routes SET value=? WHERE id=?',(json.dumps(r,ensure_ascii=False),key))
        db.execute("INSERT INTO settings(key,value) VALUES ('route_counter',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(json.dumps(next_number),))

    def routes(self):
        with self.lock,self.connect() as db:
            if any(not json.loads(v).get('number') or not json.loads(v).get('color') for v, in db.execute('SELECT value FROM commercial_routes')):
                db.execute('BEGIN IMMEDIATE');self.route_identity(db)
            return [dict(json.loads(value),id=key) for key,value in db.execute('SELECT id,value FROM commercial_routes ORDER BY rowid')]

    def save_route(self, body):
        route_id=body.get('id')
        center=body.get('center',{})
        label=str(center.get('label','')).strip()[:120]
        lat,lon,radius=float(center.get('lat',float('nan'))),float(center.get('lon',float('nan'))),float(body.get('radius',0))
        if not label or not all(math.isfinite(v) for v in (lat,lon,radius)) or not (-90<=lat<=90 and -180<=lon<=180 and 1<=radius<=1500):
            raise ValueError('Informe nome do centro, coordenadas e raio válidos.')
        keys=body.get('keys',[]);manual=body.get('manual_keys',[])
        if not isinstance(keys,list) or not isinstance(manual,list) or any(not isinstance(k,str) for k in keys+manual):raise ValueError('Seleção de municípios inválida.')
        keys=list(dict.fromkeys(keys));manual=set(manual)&set(keys)
        if not keys:raise ValueError('Escolha pelo menos um município antes de salvar.')
        with self.lock,self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            self.route_identity(db)
            records={k:json.loads(v) for k,v in db.execute('SELECT id,value FROM commercial_routes')}
            old=records.get(route_id)
            if route_id and not old:raise ValueError('Rota removida. Reabra o mapa para atualizar.')
            if old and body.get('revision')!=old.get('revision',1):raise ValueError('Esta rota foi alterada. Reabra a rota antes de salvar.')
            municipalities={k:{**json.loads(a),**json.loads(b)} for k,a,b in db.execute('SELECT key,source,overrides FROM municipalities')}
            if any(k not in municipalities or municipalities[k].get('UF')!='RS' for k in keys):raise ValueError('Selecione somente municípios cadastrados do RS.')
            occupied={k for rid,r in records.items() if rid!=route_id for k in r['keys']}
            existing_keys=set(old['keys']) if old else set()
            excluded=[k for k in keys if k in occupied and k not in manual and k not in existing_keys]
            keys=[k for k in keys if k not in excluded]
            if not keys:raise ValueError('Todos os escolhidos já pertencem a outra rota. Use + Adicionar à rota para uma inclusão manual.')
            stamp=datetime.now().astimezone();created=old['created_at'] if old else stamp.isoformat(timespec='seconds')
            name=label+' · '+format(radius,'g')+' km · '+datetime.fromisoformat(created).strftime('%d/%m/%Y')
            other_names={r['name'] for rid,r in records.items() if rid!=route_id};base=name;n=2
            while name in other_names:name=base+' ('+str(n)+')';n+=1
            route={'number':old.get('number') if old else None,'color':old.get('color') if old else None,'name':name,'uf':'RS','center':{'lat':lat,'lon':lon,'label':label},'radius':radius,'created_at':created,'updated_at':stamp.isoformat(timespec='seconds'),'revision':old.get('revision',1)+1 if old else 1,'owner':old.get('owner') if old else None,'keys':keys,'manual_keys':[k for k in keys if k in manual]}
            route_id=route_id or secrets.token_hex(12)
            db.execute('INSERT INTO commercial_routes(id,value) VALUES (?,?) ON CONFLICT(id) DO UPDATE SET value=excluded.value',(route_id,json.dumps(route,ensure_ascii=False)))
            self.route_identity(db)
            route=json.loads(db.execute('SELECT value FROM commercial_routes WHERE id=?',(route_id,)).fetchone()[0])
        return {'route':dict(route,id=route_id),'excluded_keys':excluded}

    def delete_route(self, body):
        with self.lock,self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row=db.execute('SELECT value FROM commercial_routes WHERE id=?',(body.get('id'),)).fetchone()
            if not row:raise ValueError('Rota não encontrada.')
            if body.get('revision')!=json.loads(row[0]).get('revision',1):raise ValueError('Esta rota foi alterada. Atualize o mapa antes de remover.')
            db.execute('DELETE FROM commercial_routes WHERE id=?',(body['id'],))
        return {'ok':True}

    @contextmanager
    def connect(self):
        # sqlite3's transaction context does not close its connection.
        db=sqlite3.connect(self.db_path,timeout=15)
        try:
            with db:
                yield db
        finally:
            db.close()

    def relocate_paths(self,value):
        """Resolve arquivos gerenciados copiados com dados, sem modificar o banco."""
        if isinstance(value,dict):return {k:self.relocate_paths(v) for k,v in value.items()}
        if isinstance(value,list):return [self.relocate_paths(v) for v in value]
        if not isinstance(value,str) or not re.search(r'[\\/]',value):return value
        if re.match(r'^[a-zA-Z]+://',value):return value
        parts=re.split(r'[\\/]+',value)
        indices=[i for i,p in enumerate(parts) if p.lower()=='dados']
        for i in reversed(indices):
            tail=parts[i+1:]
            if not tail or any(p in ('..','.') for p in tail):continue
            candidate=self.folder.joinpath(*tail).resolve()
            if candidate.is_relative_to(self.folder.resolve()) and candidate.exists():return str(candidate)
        return value

    def setting(self, key, default=None):
        with self.connect() as db:
            row = db.execute('SELECT value FROM settings WHERE key=?', (key,)).fetchone()
        value=json.loads(row[0]) if row else default
        return self.relocate_paths(value) if key in ('roots','catalogue','tse_states') else value

    def set_setting(self, key, value):
        with self.lock,self.connect() as db:
            db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)', (key,json.dumps(value,ensure_ascii=False)))

    def import_xlsx(self, content):
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        records = []
        for sheet in workbook:
            iterator = iter(sheet.values)
            header = next(iterator, ())
            if 'Município' not in header or 'UF' not in header:
                continue
            for values in iterator:
                record = {str(k): (v.isoformat() if hasattr(v,'isoformat') else v) for k,v in zip(header,values) if k is not None}
                if not record.get('Município') or not record.get('UF'):
                    continue
                key = str(record.get('Código Único') or f"{record['UF']}-{norm(record['Município'])}")
                records.append((key, record))
        if not records:
            raise ValueError('Nenhuma aba com as colunas Município e UF foi encontrada.')
        if len({k for k,_ in records}) != len(records):
            raise ValueError('Há identificadores de município repetidos. Confira a planilha antes de importar.')
        inserted = updated = conflicts = 0
        with self.lock, self.connect() as db:
            geography={}
            for existing_key,existing_source in db.execute('SELECT key,source FROM municipalities'):
                existing_source=json.loads(existing_source)
                geography.setdefault((str(existing_source.get('UF','')).upper(),norm(existing_source.get('Município'))),[]).append(existing_key)
            for key, source in records:
                old = db.execute('SELECT source,overrides FROM municipalities WHERE key=?', (key,)).fetchone()
                if not old:
                    matches=geography.get((str(source.get('UF','')).upper(),norm(source.get('Município'))),[])
                    if len(matches)>1:raise ValueError('Mais de um cadastro para este município. Revise antes de atualizar.')
                    if matches:
                        key=matches[0]
                        old=db.execute('SELECT source,overrides FROM municipalities WHERE key=?',(key,)).fetchone()
                if old and (str(json.loads(old[0]).get('UF','')).upper(),norm(json.loads(old[0]).get('Município'))) != (str(source.get('UF','')).upper(),norm(source.get('Município'))):raise ValueError('Código Único usado por municípios diferentes. Confira a planilha.')
                if old:
                    overrides = json.loads(old[1])
                    conflicts += sum(1 for field,value in overrides.items() if source.get(field) != value)
                    db.execute('UPDATE municipalities SET source=? WHERE key=?', (json.dumps(source,ensure_ascii=False),key))
                    updated += 1
                else:
                    db.execute('INSERT INTO municipalities (key,source) VALUES (?,?)', (key,json.dumps(source,ensure_ascii=False)))
                    geography.setdefault((str(source.get('UF','')).upper(),norm(source.get('Município'))),[]).append(key)
                    inserted += 1
        workbook.close()
        municipal_media.ensure_codes(self)
        return {'inserted':inserted,'updated':updated,'preserved_edits':conflicts}

    def items(self,only_key=None):
        with self.connect() as db:
            args=() if only_key is None else (only_key,)
            rows = db.execute('SELECT m.key,m.source,m.overrides,m.local,c.code FROM municipalities m LEFT JOIN municipality_codes c ON c.key=m.key'+(' WHERE m.key=?' if only_key is not None else ''),args).fetchall()
            projects={key:self.relocate_paths(json.loads(value)) for key,value in db.execute('SELECT key,value FROM projects'+(' WHERE key=?' if only_key is not None else ''),args)}
            activity={};send_states=whatsapp_send126.states(db);locals_by_key={r[0]:json.loads(r[3]) for r in rows}
            for event_id,key,timestamp,kind,channel,target,recipient,pdf_name in db.execute('SELECT id,municipality_key,timestamp,kind,channel,target,recipient,pdf_name FROM interactions'+(' WHERE municipality_key=?' if only_key is not None else '')+(' AND ' if only_key is not None else ' WHERE ')+"NOT EXISTS (SELECT 1 FROM interaction_corrections c WHERE c.id=interactions.id AND c.active=0) ORDER BY timestamp,rowid",args):
                if kind=='sent' and channel=='whatsapp' and not whatsapp_send126.allowed(locals_by_key.get(key,{}),target,recipient,send_states.get(event_id)):continue
                state=activity.setdefault(key,{'send_count':0,'channels':set(),'sent_channels':set(),'last_send':None,'last_interaction':None})
                summary={'timestamp':timestamp,'kind':kind,'channel':channel,'target':target,'recipient':recipient,'pdf_name':pdf_name}
                state['last_interaction']=summary;state['channels'].add(channel+'_'+target)
                if kind=='sent':
                    state['send_count']+=1;state['last_send']=summary;state['sent_channels'].add(channel+'_'+target)
                    send_activity130.add(state,summary)
                    if pdf_name:state['pdf_send_count']=state.get('pdf_send_count',0)+1;state['last_pdf_send']=summary
        result = []
        for key,source,overrides,local,municipal_code in rows:
            source,overrides,local = json.loads(source),json.loads(overrides),json.loads(local)
            fields = {**source,**overrides}
            fields.setdefault('Diretor da Câmara (Nome)','')
            fields.setdefault('Sec. Executivo / Atendimento (Nome)','')
            uf,city = str(fields.get('UF','')).upper(),fields.get('Município','')
            allowed = EXCEPTIONS.get((uf,norm(city)), UF_DDD.get(uf, []))
            primary = ddds(fields.get('Telefone Planilha Original'))
            collected = ddds(fields.get('Telefones Coletados no Site'))
            candidates = [d for d in primary + collected if d in allowed]
            geographic = str(local.get('ddd',''))
            if not geographic and len(set(candidates)) == 1:
                geographic = candidates[0]
            warning = bool([d for d in primary + collected if d not in allowed])
            conflicts = [f for f,v in overrides.items() if source.get(f) != v]
            project = projects.get(key,{})
            if project.get('pdf') and (project.get('catalogue_version')!=1 or project.get('layout_version')!='1.33.1'):project['dirty']=True
            pairs = project.get('pairs',[])
            status = 'Sem fotos' if not pairs else 'Em restauração'
            if pairs and all(p.get('restored') and Path(p['restored']).is_file() for p in pairs if p.get('selected',True)) and any(p.get('selected',True) for p in pairs):
                status = 'Pronto para PDF'
            if project.get('pdf') and not project.get('dirty',False): status = 'PDF pronto'
            interactions=activity.get(key,{'send_count':0,'channels':set(),'sent_channels':set(),'last_send':None,'last_interaction':None})
            contacted=bool(interactions['last_interaction'] or local.get('last_contact') or local.get('commercial','Não contatado')!='Não contatado')
            result.append({'send_days':send_activity130.output(interactions),'pdf_send_count':interactions.get('pdf_send_count',0),'last_pdf_send':interactions.get('last_pdf_send'),'send_count':interactions['send_count'],'sent_channels':sorted(interactions['sent_channels']),'interaction_channels':sorted(interactions['channels']),'last_send':interactions['last_send'],'last_interaction':interactions['last_interaction'],'contacted':contacted,'key':key,'code':municipal_code,'fields':fields,'source':source,'local':local,'overrides':list(overrides),'conflicts':conflicts,'ddd':geographic,'ddd_inferred':not bool(local.get('ddd')) and bool(geographic),'ddd_warning':warning,'allowed_ddds':allowed,'stage':status})
        return sorted(result, key=lambda x:(x['fields'].get('UF',''),x['fields'].get('Município','')))

    def item(self,key):
        item = next(iter(self.items(key)),None)
        if item is None: raise ValueError('Município não encontrado.')
        return item

    def save_item(self,key,fields,local,manual=False):
        with self.lock, self.connect() as db:
            row = db.execute('SELECT source,overrides,local FROM municipalities WHERE key=?',(key,)).fetchone()
            if not row: raise ValueError('Município não encontrado.')
            source, overrides, metadata = map(json.loads,row)
            before_fields={**source,**overrides};before_local=dict(metadata)
            fields=dict(fields)
            if 'Presidente - Email' in fields:fields['Presidente - Email']=editing_activity.single_email(fields['Presidente - Email'])
            previous_website={**source,**overrides}.get('Website Câmara','')
            changes_pdf=any(field in ('Município','UF') and {**source,**overrides}.get(field)!=value for field,value in fields.items())
            for field,value in fields.items():
                if field not in source:
                    if field not in ('Diretor da Câmara (Nome)','Sec. Executivo / Atendimento (Nome)','Website Câmara','URL Mesa Diretora','Telefone Planilha Original','Telefones Coletados no Site','WhatsApp(s) da Câmara (Botões/Links)','Email Geral da Câmara','Endereço da Câmara','URL Contato','Presidente Atual (Nome)','Presidente - Telefone / Celular','Presidente - WhatsApp','Presidente - Email','Presidente - URL Perfil','Presidente - Partido','Presidente - Período Mandato'):raise ValueError('Campo desconhecido: '+field)
                    source[field]=''
                if source.get(field)==value: overrides.pop(field,None)
                else: overrides[field]=value
            if 'Website Câmara' in fields and previous_website!=fields['Website Câmara']:metadata['website_manual_at']=datetime.now().astimezone().isoformat(timespec='seconds')
            metadata.update(local)
            if manual and (any(before_fields.get(k,'')!=v for k,v in fields.items()) or any(before_local.get(k)!=v for k,v in local.items())):editing_activity.mark(metadata)
            metadata['updated_at']=now()
            if local.get('checked'): metadata['checked_at']=now()
            db.execute('UPDATE municipalities SET source=?,overrides=?,local=? WHERE key=?',(json.dumps(source,ensure_ascii=False),json.dumps(overrides,ensure_ascii=False),json.dumps(metadata,ensure_ascii=False),key))
        if changes_pdf:self.invalidate(key)

    def interaction_history(self,key):
        with self.connect() as db:
            municipality=db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()
            if not municipality:raise ValueError('Município não encontrado.')
            local=json.loads(municipality[0]);send_states=whatsapp_send126.states(db)
            rows=db.execute('SELECT id,timestamp,kind,channel,target,recipient,pdf_path,pdf_name,notes FROM interactions WHERE municipality_key=? ORDER BY timestamp DESC,rowid DESC',(key,)).fetchall()
        fields=('id','timestamp','kind','channel','target','recipient','pdf_path','pdf_name','notes')
        with self.connect() as db:
            corrections={i:{'active':bool(a),'corrections':json.loads(audit)} for i,a,audit in db.execute('SELECT id,active,audit FROM interaction_corrections')}
            versions=dict(db.execute('SELECT id,sha256 FROM interaction_files'))
        result=[]
        for row in rows:
            event=dict(zip(fields,row),pdf_sha256=versions.get(row[0]),**corrections.get(row[0],{'active':True,'corrections':[]}))
            if event['channel']=='whatsapp' and event['kind']=='sent':
                event['verification_state']=send_states.get(event['id']) or ('confirmed' if whatsapp_send126.confirmed(local,event['target'],event['recipient']) else 'pending')
                event['active']=event['active'] and whatsapp_send126.allowed(local,event['target'],event['recipient'],send_states.get(event['id']))
            result.append(event)
        return result

    def check_phone(self,key,number):
        item=self.item(key);digits=normalize_phone(number);national=digits[2:]
        return {'number':digits,'display':'('+national[:2]+') '+national[2:-4]+'-'+national[-4:],
                'ddd':national[:2],'type':'Celular' if len(national)==11 else 'Fixo',
                'ddd_warning':national[:2] not in item['allowed_ddds'],'allowed_ddds':item['allowed_ddds']}

    def add_phone(self,key,payload):
        target=payload.get('target')
        if target not in ('camara','presidente'):raise ValueError('Escolha Câmara ou presidente.')
        checked=self.check_phone(key,payload.get('number'))
        if checked['ddd_warning'] and payload.get('confirm_ddd') is not True:
            raise ValueError('DDD diferente da referência municipal. Confira e marque a exceção antes de salvar.')
        fields=('Telefone Planilha Original','WhatsApp(s) da Câmara (Botões/Links)') if target=='camara' else ('Presidente - Telefone / Celular','Presidente - WhatsApp')
        with self.lock,self.connect() as db:
            row=db.execute('SELECT source,overrides,local FROM municipalities WHERE key=?',(key,)).fetchone()
            if not row:raise ValueError('Município não encontrado.')
            source,overrides,metadata=map(json.loads,row)
            before={**source,**overrides};old_verified=metadata.get('phone_status',{}).get(target,{}).get(checked['number'],{}).get('whatsapp_verified',False)
            for field in fields:
                source.setdefault(field,'')
                previous={**source,**overrides}.get(field,'')
                if checked['number'] not in phone_numbers(previous):
                    overrides[field]=(str(previous).strip()+'; ' if useful(previous) else '')+checked['display']
            status=metadata.setdefault('phone_status',{}).setdefault(target,{})
            status[checked['number']]={'whatsapp_verified':payload.get('whatsapp_verified') is True,'checked_at':now(),'ddd_exception':checked['ddd_warning']}
            if any(before.get(f,'')!={**source,**overrides}.get(f,'') for f in fields) or old_verified!=(payload.get('whatsapp_verified') is True):editing_activity.mark(metadata)
            metadata['updated_at']=now()
            db.execute('UPDATE municipalities SET source=?,overrides=?,local=? WHERE key=?',(json.dumps(source,ensure_ascii=False),json.dumps(overrides,ensure_ascii=False),json.dumps(metadata,ensure_ascii=False),key))
        if old_verified and payload.get('whatsapp_verified') is not True:whatsapp_send126.revoke(self,key,target,checked['number'])
        return {'ok':True,**checked}

    def start_capture(self,key,url):
        item=self.item(key);parsed=urlparse(str(url or ''))
        if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password:raise ValueError('Informe um site válido para abrir a captura.')
        ticket=secrets.token_hex(24)
        with self.lock:
            self.capture_tickets={k:v for k,v in self.capture_tickets.items() if v['expires']>time.time()}
            self.capture_tickets[ticket]={'key':key,'url':url,'operator':operators131.actor(),'expires':time.time()+600}
        return {'ticket':ticket,'municipality':item['fields']['Município'],'uf':item['fields']['UF']}

    def connect_capture(self,ticket,tab_id=None):
        if type(tab_id) is not int or tab_id<0:raise ValueError('Atualize a extensão e abra a captura novamente pelo MVP.')
        with self.lock:
            binding=self.capture_tickets.pop(str(ticket),None)
            if not binding or binding['expires']<time.time():raise ValueError('Vínculo expirado ou já utilizado. Abra Capturar no site novamente no MVP.')
            token=secrets.token_hex(32);self.capture_sessions[token]={**binding,'expires':time.time()+8*3600,'pages':{str(tab_id):capture_routes.canonical(binding['url'])} if tab_id is not None else {}}
        item=self.item(binding['key'])
        return {'token':token,'key':binding['key'],'code':item['code'],'url':binding['url'],'scope':binding.get('scope','contacts'),'expected_number':binding.get('expected_number'),'municipality':item['fields']['Município'],'uf':item['fields']['UF']}

    def capture_binding(self,token):
        with self.lock:binding=self.capture_sessions.get(str(token))
        if not binding or binding['expires']<time.time():raise ValueError('Captura desconectada. Abra Capturar no site no MVP novamente.')
        return binding

    def capture_preview(self,token,body):
        binding=capture_routes.check(self,token,body,'contacts')
        if binding.get('scope')=='crest':raise ValueError('Aba de brasão: use Salvar como brasão na imagem.')
        item=self.item(binding['key'])
        fields={'name':'Presidente Atual (Nome)','phone_president':'Presidente - Telefone / Celular','phone_chamber':'Telefone Planilha Original','email_chamber':'Email Geral da Câmara','email_president':'Presidente - Email','whatsapp_chamber':'WhatsApp(s) da Câmara (Botões/Links)','whatsapp_president':'Presidente - WhatsApp'}
        kind=body.get('kind');field=fields.get(kind)
        if not field:raise ValueError('Campo de captura inválido.')
        value=str(body.get('value','')).strip();url=str(body.get('url',''));parsed=urlparse(url)
        if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password:raise ValueError('Página de origem inválida.')
        if kind.startswith('whatsapp_'):
            return self.preview_whatsapp_capture(binding,item,body,field,url,parsed)
        if not value or len(value)>500:raise ValueError('Selecione um único dado com até 500 caracteres.')
        checked={}
        if kind.startswith('phone_'):
            if re.match(r'^tel\s*:',value,re.I):value=unquote(re.sub(r'^tel\s*:', '',value,flags=re.I).split('?',1)[0].split(';',1)[0]).strip()
            checked=self.check_phone(binding['key'],value);value=checked['display']
        elif kind.startswith('email_'):
            value=unquote(value).strip()
            value=re.sub(r'^(?:mailto|e-?mail)\s*:\s*','',value,flags=re.I).strip()
            value=value.split('?',1)[0].split('#',1)[0].strip()
            if not re.fullmatch(r"[^\s@?&#;,<>:/\\]+@[^\s@?&#;,<>:/\\]+\.[^\s@?&#;,<>:/\\]+",value):raise ValueError('Selecione apenas um endereço de e-mail válido.')
        elif len(value)>200 or '\n' in value:raise ValueError('Selecione somente o nome do presidente.')
        old=str(item['fields'].get(field,'') or '')
        return {'key':binding['key'],'municipality':item['fields']['Município'],'uf':item['fields']['UF'],'field':field,'kind':kind,'value':value,'previous':old,'url':url,'source_changed':parsed.hostname!=urlparse(binding['url']).hostname,'replace_required':kind=='name' and useful(old) and old!=value,**checked}

    def preview_whatsapp_capture(self,binding,item,body,field,url,parsed):
        link=parse_link(body.get('value'),normalize_phone)
        target='camara' if body['kind']=='whatsapp_chamber' else 'presidente'
        phone_field='Telefone Planilha Original' if target=='camara' else 'Presidente - Telefone / Celular'
        numbers=list(dict.fromkeys(phone_numbers(item['fields'].get(phone_field,''))+phone_numbers(item['fields'].get(field,''))))
        selected=normalize_phone(body['selected_number']) if body.get('selected_number') else ''
        if selected and selected not in numbers:raise ValueError('Escolha um telefone já cadastrado para este contato.')
        if link['link_number'] and selected and selected!=link['link_number']:
            raise ValueError('O número escolhido difere do número do link. Use o número presente no link ou confira outro endereço.')
        number=link['link_number'] or selected
        checked=self.check_phone(binding['key'],number) if number else {}
        previous_link=item['local'].get('whatsapp_links',{}).get(target,{}).get(number,{}).get('url','')
        replace=bool(previous_link and previous_link!=link['link_url'])
        return {'key':binding['key'],'municipality':item['fields']['Município'],'uf':item['fields']['UF'],
                'field':field,'kind':body['kind'],'value':link['link_url'],'previous':str(item['fields'].get(field,'') or ''),
                'url':url,'target':target,'selected_number':selected,'numbers':numbers,'previous_link':previous_link,
                'source_changed':parsed.hostname!=urlparse(binding['url']).hostname,'replace_required':False,
                'link_replace_required':replace,'requires_review':link['opaque'] or replace or checked.get('ddd_warning',False),**link,**checked}

    def save_whatsapp_capture(self,preview,body):
        key=preview['key'];target=preview['target'];number=preview.get('number','')
        if preview['opaque'] and body.get('confirm_opaque') is not True:
            raise ValueError('Link sem número: confira o contato e escolha a associação antes de salvar.')
        if preview['link_replace_required']:
            if body.get('confirm_link_replace') is not True or body.get('expected_link')!=preview['previous_link']:
                raise ValueError('O endereço deste WhatsApp mudou. Confira e confirme a substituição.')
        elif 'expected_link' in body and body['expected_link']!=preview['previous_link']:
            raise ValueError('O endereço deste WhatsApp mudou. Confira novamente.')
        if number:
            verified=bool(self.item(key)['local'].get('phone_status',{}).get(target,{}).get(number,{}).get('whatsapp_verified'))
            self.add_phone(key,{'target':target,'number':number,'confirm_ddd':body.get('confirm_ddd'),'whatsapp_verified':verified})
        with self.connect() as db:
            row=db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone();local=json.loads(row[0])
            links=local.setdefault('whatsapp_links',{}).setdefault(target,{})
            if links.get(number,{}).get('url')!=preview['link_url']:editing_activity.mark(local)
            links[number]={'url':preview['link_url'],'source_url':preview['url'],'captured_at':now(),'opaque':preview['opaque']}
            local['updated_at']=now()
            db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))

    def save_capture(self,token,body):
        with self.lock:
            preview=self.capture_preview(token,body);key=preview['key'];kind=preview['kind'];field=preview['field'];value=preview['value']
            direct=body.get('direct') is True
            if not direct and preview['previous']!=str(body.get('expected_previous','')):raise ValueError('Este campo mudou no MVP. Confira novamente antes de salvar.')
            if not direct and preview['source_changed'] and body.get('confirm_source') is not True:raise ValueError('Confirme a captura de um site diferente do site inicial.')
            if not direct and preview['replace_required'] and body.get('confirm_replace') is not True:raise ValueError('Confirme a substituição do nome atual.')
            if kind.startswith('whatsapp_'):
                self.save_whatsapp_capture(preview,body)
            elif kind.startswith('phone_'):
                number=preview['number'];item=self.item(key);target='camara' if kind=='phone_chamber' else 'presidente'
                verified=bool(item['local'].get('phone_status',{}).get(target,{}).get(number,{}).get('whatsapp_verified'))
                self.add_phone(key,{'target':target,'number':number,'confirm_ddd':body.get('confirm_ddd'),'whatsapp_verified':verified})
            else:
                old=preview['previous']
                if kind=='email_chamber':
                    addresses=re.findall(r"[^\s@?&#;,]+@[^\s@?&#;,]+\.[^\s@?&#;,]+",old)
                    value=old if value.lower() in [v.lower() for v in addresses] else (old.strip()+'; ' if useful(old) else '')+value
                self.save_item(key,{field:value},{},manual=True)
            # Audit source, separate from conversations/sends; original captured text stays visible.
            with self.connect() as db:
                row=db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone();local=json.loads(row[0]);audit=local.setdefault('captures',[])
                audit.append({'id':secrets.token_hex(12),'timestamp':datetime.now().astimezone().isoformat(timespec='seconds'),'field':field,'value':preview['value'],'selected_text':str(body.get('selected_text',body.get('value','')))[:1000],'url':preview['url']})
                db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))
            self.capture_sequence+=1;self.capture_events.append((self.capture_sequence,key));self.capture_events=self.capture_events[-200:]
            self.capture_updates.append({'sequence':self.capture_sequence,'key':key,'field':field});self.capture_updates=self.capture_updates[-200:]
            self.capture_condition.notify_all()
            return {'ok':True,'municipality':preview['municipality'],'uf':preview['uf'],'field':field,'value':preview['value'],'sequence':self.capture_sequence}

    def record_interaction(self,key,payload):
        event_id=payload.get('id','')
        if not re.fullmatch(r'[a-f0-9]{24}',event_id):raise ValueError('Identificador de registro inválido.')
        kind=payload.get('kind');channel=payload.get('channel');target=payload.get('target');recipient=str(payload.get('recipient','')).strip()
        if kind not in ('sent','contact'):raise ValueError('Tipo de registro inválido.')
        if channel not in ('whatsapp','email','other'):raise ValueError('Canal inválido.')
        if target not in ('camara','presidente','outro'):raise ValueError('Destinatário inválido.')
        if kind=='sent' and channel=='other':raise ValueError('Escolha WhatsApp ou e-mail para registrar um envio.')
        if channel=='whatsapp':
            saved_links=self.item(key)['local'].get('whatsapp_links',{}).get(target,{})
            opaque=any(v.get('opaque') and v.get('url')==recipient for v in saved_links.values())
            if not opaque:
                field='WhatsApp(s) da Câmara (Botões/Links)' if target=='camara' else 'Presidente - WhatsApp' if target=='presidente' else None
                if field and str(self.item(key)['fields'].get(field,'')).strip()==recipient:
                    try:opaque=parse_link(recipient,normalize_phone)['opaque']
                    except ValueError:pass
            if not opaque:
                recipient=re.sub(r'\D','',recipient)
                if len(recipient) in (10,11):recipient='55'+recipient
                if not re.fullmatch(r'55\d{10,11}',recipient):raise ValueError('Informe um número válido com DDD ou um link sem número salvo neste contato.')
        if channel=='email' and not re.fullmatch(r'[^\s@?&#]+@[^\s@?&#]+\.[^\s@?&#]+',recipient):raise ValueError('Informe um e-mail válido.')
        notes=str(payload.get('notes','')).strip()[:10000]
        with self.lock,self.connect() as db:
            row=db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()
            if not row:raise ValueError('Município não encontrado.')
            prior=db.execute('SELECT municipality_key,kind,channel,target,recipient FROM interactions WHERE id=?',(event_id,)).fetchone()
            if prior:
                if prior!=(key,kind,channel,target,recipient):raise ValueError('Registro de envio em conflito.')
                return {'ok':True,'id':event_id,'duplicate':True,'verification_state':whatsapp_send126.states(db).get(event_id)}
            pdf_path=pdf_name=None
            if payload.get('attach_pdf'):
                project_row=db.execute('SELECT value FROM projects WHERE key=?',(key,)).fetchone()
                project=json.loads(project_row[0]) if project_row else {}
                if not project.get('pdf') or project.get('dirty') or (project.get('catalogue_version')!=1 or project.get('layout_version')!='1.33.1') or not Path(project['pdf']).is_file():raise ValueError('Gere um PDF final atualizado antes de registrar o envio desse arquivo.')
                if payload.get('expected_pdf') and payload['expected_pdf']!=project['pdf']:raise ValueError('O PDF mudou depois da seleção de destinatários. Confira o novo arquivo antes de registrar.')
                pdf_path=project['pdf'];pdf_name=Path(pdf_path).name
            timestamp=datetime.now().astimezone().isoformat(timespec='seconds')
            db.execute('INSERT INTO interactions VALUES (?,?,?,?,?,?,?,?,?,?)',(event_id,key,timestamp,kind,channel,target,recipient,pdf_path,pdf_name,notes))
            local=json.loads(row[0]);verification_state=None
            if channel=='whatsapp' and kind=='sent':
                verification_state='confirmed' if whatsapp_send126.confirmed(local,target,recipient) else 'pending' if re.fullmatch(r'55\d{10,11}',recipient) else 'opaque'
                db.execute('INSERT INTO whatsapp_send_checks VALUES (?,?)',(event_id,verification_state))
            if pdf_path:
                import hashlib
                db.execute('INSERT INTO interaction_files VALUES (?,?)',(event_id,hashlib.sha256(Path(pdf_path).read_bytes()).hexdigest()))
            if verification_state in ('pending','opaque'):return {'ok':True,'id':event_id,'timestamp':timestamp,'verification_state':verification_state}
            local.setdefault('interaction_previous',{'last_contact':local.get('last_contact'),'commercial':local.get('commercial')});local['last_contact']=timestamp[:10];local['updated_at']=timestamp
            if local.get('commercial','Não contatado')=='Não contatado':local['commercial']='Contatado'
            db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))
        return {'ok':True,'id':event_id,'timestamp':timestamp,'verification_state':verification_state}

    def correct_interaction(self,key,body):
        active=body.get('active')
        if type(active) is not bool:raise ValueError('Informe a situação do registro.')
        with self.lock,self.connect() as db:
            row=db.execute('SELECT timestamp FROM interactions WHERE id=? AND municipality_key=?',(body.get('id'),key)).fetchone()
            if not row:raise ValueError('Registro não encontrado.')
            event=db.execute('SELECT channel,kind,target,recipient FROM interactions WHERE id=?',(body['id'],)).fetchone()
            local=json.loads(db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()[0])
            if active and event[0]=='whatsapp' and event[1]=='sent' and not whatsapp_send126.allowed(local,event[2],event[3],whatsapp_send126.states(db).get(body['id'])):raise ValueError('Confirme primeiro se este número tem WhatsApp.')
            if not active and whatsapp_send126.states(db).get(body['id'])=='pending':db.execute("UPDATE whatsapp_send_checks SET state='rejected' WHERE id=?",(body['id'],))
            old=db.execute('SELECT active,audit FROM interaction_corrections WHERE id=?',(body['id'],)).fetchone()
            if old and bool(old[0])==active:return {'ok':True}
            audit=json.loads(old[1]) if old else [];audit.append({'active':active,'timestamp':now()})
            db.execute('INSERT OR REPLACE INTO interaction_corrections VALUES (?,?,?)',(body['id'],int(active),json.dumps(audit)))
            candidates=db.execute('SELECT id,timestamp,channel,kind,target,recipient FROM interactions i WHERE municipality_key=? AND NOT EXISTS (SELECT 1 FROM interaction_corrections c WHERE c.id=i.id AND c.active=0) ORDER BY timestamp DESC,rowid DESC',(key,)).fetchall()
            checks=whatsapp_send126.states(db)
            remaining=next(((e[1],) for e in candidates if e[2]!='whatsapp' or e[3]!='sent' or whatsapp_send126.allowed(local,e[4],e[5],checks.get(e[0]))),None)
            if active:
                local=json.loads(db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()[0]);local.setdefault('interaction_previous',{'last_contact':local.get('last_contact'),'commercial':local.get('commercial')})
                local['last_contact']=remaining[0][:10] if remaining else row[0][:10]
                if local.get('commercial','Não contatado')=='Não contatado':local['commercial']='Contatado'
                db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))
            if not active and remaining:
                local=json.loads(db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()[0])
                recorded_dates={r[0][:10] for r in db.execute('SELECT timestamp FROM interactions WHERE municipality_key=?',(key,))}
                if local.get('last_contact') in recorded_dates:local['last_contact']=remaining[0][:10]
                db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))
            if not remaining:
                local=json.loads(db.execute('SELECT local FROM municipalities WHERE key=?',(key,)).fetchone()[0]);previous=local.pop('interaction_previous',{})
                if local.get('last_contact')==row[0][:10]:local['last_contact']=previous.get('last_contact')
                if local.get('commercial')=='Contatado':local['commercial']=previous.get('commercial') or 'Não contatado'
                db.execute('UPDATE municipalities SET local=? WHERE key=?',(json.dumps(local,ensure_ascii=False),key))
        return {'ok':True}

    def project(self,key):
        with self.connect() as db:
            row = db.execute('SELECT value FROM projects WHERE key=?',(key,)).fetchone()
        project=self.relocate_paths(json.loads(row[0])) if row else {'pairs':[],'restored_files':[],'year':2004,'period':'2005–2008','dirty':True}
        if project.get('pdf') and (project.get('catalogue_version')!=1 or project.get('layout_version')!='1.33.1'):project['dirty']=True
        return project

    def save_project(self,key,value):
        with self.connect() as db:
            db.execute('INSERT OR REPLACE INTO projects VALUES (?,?)',(key,json.dumps(value,ensure_ascii=False)))

    def invalidate(self,key):
        p=self.project(key);p['dirty']=True;self.save_project(key,p)

    def root_settings(self):
        return self.setting('roots',{})

    def tse_import(self,uf,content,name,progress=None):
        uf=str(uf).upper()
        if uf!='BR' and (uf not in UF_DDD or uf=='DF'):raise ValueError('Selecione um estado ou Brasil inteiro.')
        if progress:progress('Lendo a consulta de candidatos de 2004…',5)
        candidates=read_candidates(content,name,uf,progress)
        if progress:progress('Validando municípios e preservando cadastros existentes…',40)
        states=self.setting('tse_states',{})
        existing={}
        with self.connect() as db:
            for key,source in db.execute('SELECT key,source FROM municipalities'):
                source=json.loads(source);existing.setdefault((str(source.get('UF','')).upper(),norm(source.get('Município'))),[]).append(key)
        grouped={}
        for c in candidates:grouped.setdefault((c['uf'],norm(c['city'])),[]).append(c)
        if any(len(existing.get(geo,[]))>1 for geo in grouped):raise ValueError('Há municípios duplicados. Revise antes de importar.')
        columns=['ID','Código Único','Código IBGE','Arquivado','Status Extração','Município','UF','Website Câmara','Validação Website','Telefone Planilha Original','Presidente Atual (Nome)','Presidente - Partido','Presidente - Período Mandato','Presidente - Telefone / Celular','Presidente - WhatsApp','Presidente - Email','Presidente - URL Perfil','WhatsApp(s) da Câmara (Botões/Links)','Telefones Coletados no Site','Email Geral da Câmara','Endereço da Câmara','Mesa Diretora Completa','Desenvolvedor / Plataforma','URL Mesa Diretora','URL Contato','Nível de Confiança','Notas da Extração','Fonte / Páginas Analisadas','Data da Extração','Mensagem de Erro']
        totals={state:{'new':0,'matched':0,'candidates':[]} for state,_ in grouped}
        # One transaction for every UF: validation errors cannot leave a partial national import.
        with self.lock,self.connect() as db:
            for (state,city),people in grouped.items():
                keys=existing.get((state,city),[])
                if keys:key=keys[0];totals[state]['matched']+=1
                else:
                    key=f'{state}-{city}'
                    if db.execute('SELECT 1 FROM municipalities WHERE key=?',(key,)).fetchone():raise ValueError('Identificador em conflito: '+key)
                    source=dict.fromkeys(columns)
                    source.update({'ID':key,'Código Único':key,'UF':state,'Município':people[0]['city'],'Arquivado':'NÃO','Status Extração':'Pendente','Notas da Extração':'Município cadastrado a partir do TSE. Contatos ainda pendentes.'})
                    db.execute('INSERT INTO municipalities (key,source) VALUES (?,?)',(key,json.dumps(source,ensure_ascii=False)));totals[state]['new']+=1
                for candidate in people:candidate['key']=key
                totals[state]['candidates'].extend(people)
            for state,total in totals.items():
                previous=states.get(state,{})
                merged={(c.get('key'),c['id']):c for c in previous.get('candidates',[])}
                merged.update({(c['key'],c['id']):c for c in total['candidates']})
                previous.update({'uf':state,'year':2004,'filename':Path(name).name,'source_scope':uf,'imported_at':now(),'candidates':list(merged.values()),'new_municipalities':total['new'],'matched_municipalities':total['matched'],'summary':None})
                states[state]=previous
            db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',('tse_states',json.dumps(states,ensure_ascii=False)))
        folder=self.folder/'tse'/uf;folder.mkdir(parents=True,exist_ok=True)
        (folder/('consulta'+Path(name).suffix.lower())).write_bytes(content)
        errors=[]
        for i,state in enumerate(sorted(totals)):
            if progress:progress('Vinculando fotos de '+state+'…',50+int(i/max(1,len(totals))*45))
            root=states[state].get('photo_root')
            if root and Path(root).is_dir():
                try:self.tse_process(state,root)
                except ValueError as e:errors.append({'uf':state,'error':str(e)})
        if uf!='BR':return self.tse_public(uf)
        public=self.tse_public()
        result={'scope':'BR','filename':Path(name).name,'imported_at':now(),'elected':len(candidates),'municipalities':len(grouped),'new_municipalities':sum(x['new'] for x in totals.values()),'matched_municipalities':sum(x['matched'] for x in totals.values()),'states':{state:public[state] for state in sorted(totals)},'photo_errors':errors}
        self.set_setting('tse_national',result)
        return result

    def start_tse_import(self,uf,content,name):
        if not self.import_lock.acquire(blocking=False):raise ValueError('Já há uma importação TSE em andamento. Aguarde a conclusão.')
        job=secrets.token_hex(12)
        with self.job_lock:self.jobs[job]={'status':'running','message':'Consulta recebida. Iniciando…','percent':0}
        def progress(message,percent):
            with self.job_lock:
                self.jobs[job]['message']=message
                if percent is not None:self.jobs[job]['percent']=percent
        def work():
            try:
                result=self.tse_import(uf,content,name,progress)
                with self.job_lock:self.jobs[job].update({'status':'done','percent':100,'message':'Importação concluída.','result':result})
            except Exception as error:
                with self.job_lock:self.jobs[job].update({'status':'error','message':str(error)})
            finally:self.import_lock.release()
        context=contextvars.copy_context()
        threading.Thread(target=lambda:context.run(work),daemon=True).start()
        return {'job':job}

    def import_status(self,job):
        with self.job_lock:
            if job not in self.jobs:raise ValueError('Importação não encontrada. O programa pode ter sido reiniciado.')
            return dict(self.jobs[job])

    def tse_public(self,uf=None):
        states=self.setting('tse_states',{})
        def view(state):
            return {k:v for k,v in state.items() if k!='candidates'} | {'elected':len(state.get('candidates',[])),'municipalities':len({x['key'] for x in state.get('candidates',[])})}
        if uf:return view(states.get(uf,{}))
        return {k:view(v) for k,v in states.items()}

    def tse_photo_upload(self,uf,relative,content,batch):
        if uf not in UF_DDD or uf=='DF':raise ValueError('Estado inválido.')
        if not re.fullmatch(r'[a-f0-9]{16,40}',batch):raise ValueError('Lote de upload inválido.')
        relative=relative.replace('\\','/')
        parts=relative.split('/')
        if any(x in ('','..','.') or ':' in x for x in parts):raise ValueError('Nome de arquivo inválido.')
        if Path(relative).suffix.lower() not in IMAGE_EXT:raise ValueError('Use somente imagens na pasta de fotos.')
        folder=self.folder/'tse'/uf/'uploads'/batch
        target=folder.joinpath(*parts);target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(content)
        return {'folder':str(folder.resolve())}

    def tse_process(self,uf,path=None):
        states=self.setting('tse_states',{});state=states.get(uf)
        if not state:raise ValueError('Importe primeiro a consulta de candidatos de 2004 deste estado.')
        configured=str(path or state.get('photo_root','')).strip()
        if not configured:raise ValueError('Informe a pasta das fotos DIV ou faça o upload da pasta.')
        root=Path(configured)
        if not root.is_dir():raise ValueError('Pasta das fotos DIV não encontrada. Descompacte as fotos primeiro.')
        by_base={}
        for file in root.rglob('*'):
            if file.is_file() and file.suffix.lower() in IMAGE_EXT:by_base.setdefault(clean(file.stem),[]).append(file)
        if not by_base:raise ValueError('Nenhuma imagem encontrada na pasta das fotos DIV.')
        if not any(clean(c['expected']) in by_base for c in state['candidates']):
            raise ValueError('Nenhuma foto DIV corresponde aos eleitos de 2004 deste estado. Confira ano, UF e os nomes originais das fotos.')
        output=self.folder/'tse'/uf/'eleitos';output.mkdir(parents=True,exist_ok=True)
        grouped={}
        for candidate in state['candidates']:grouped.setdefault(candidate['key'],[]).append(candidate)
        found=missing=ambiguous=0;report=[]
        with self.lock:
            for key,candidates in grouped.items():
                project=self.project(key);old=project.get('pairs',[])
                old_ids={x.get('tse_id'):x for x in old if x.get('tse_id')}
                old_names={}
                for x in old:old_names.setdefault(x['match'],[]).append(x)
                names={}
                for c in candidates:names[norm(c['name'])]=names.get(norm(c['name']),0)+1
                pairs=[];municipal_report=[]
                for c in candidates:
                    matches=by_base.get(clean(c['expected']),[])
                    status='ENCONTRADA' if len(matches)==1 else ('AMBÍGUA' if matches else 'NÃO ENCONTRADA')
                    record={**c,'photo_status':status,'source_files':[str(x) for x in matches]}
                    municipal_report.append(record);report.append(record)
                    if len(matches)!=1:
                        if matches:ambiguous+=1
                        else:missing+=1
                        continue
                    found+=1
                    municipal_folder=output/safe_filename(clean(c['city']))/'2004'
                    sub=Path('HOMONIMOS')/c['id'] if names[norm(c['name'])]>1 else Path('.')
                    file=municipal_folder/sub/(safe_filename(clean(c['name']))+matches[0].suffix.lower())
                    file.parent.mkdir(parents=True,exist_ok=True)
                    same=file.exists() and file.read_bytes()==matches[0].read_bytes()
                    shutil.copy2(matches[0],file)
                    previous=old_ids.get(c['id'])
                    if not previous and len(old_names.get(norm(c['name']),[]))==1 and names[norm(c['name'])]==1:previous=old_names[norm(c['name'])][0]
                    relative=file.relative_to(municipal_folder)
                    pair={'path':str(file.resolve()),'name':c['name'],'match':norm(c['name']),'relative':str(relative),'identity':'/'.join([norm(x) for x in relative.parts[:-1]]+[norm(file.stem)]),'id':previous['id'] if previous else secrets.token_hex(8),'original':str(file.resolve()),'restored':previous.get('restored') if previous else None,'selected':previous.get('selected',True) if previous else True,'approved':bool(previous and same and previous.get('approved')),'tse_id':c['id']}
                    pairs.append(pair)
                project.update({'pairs':pairs,'tse_records':municipal_report,'tse_uf':uf,'original_folder':str((output/safe_filename(clean(candidates[0]['city']))/'2004').resolve()),'dirty':True})
                self.save_project(key,project)
        roots=self.root_settings();roots.setdefault(uf,{})['original']=str(output.resolve());self.set_setting('roots',roots)
        summary={'found':found,'missing':missing,'ambiguous':ambiguous,'municipalities':len(grouped),'elected':len(state['candidates'])}
        state.update({'photo_root':str(root.resolve()),'processed_at':now(),'summary':summary});states[uf]=state;self.set_setting('tse_states',states)
        report_path=self.folder/'tse'/uf/'relatorio_fotos_2004.json';report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
        return self.tse_public(uf)

    def photo_folder(self,key,override=None,kind='original'):
        item=self.item(key); uf=item['fields']['UF']; city=item['fields']['Município']
        root = Path(str(override or self.root_settings().get(uf,{}).get(kind,'')))
        if not str(override or self.root_settings().get(uf,{}).get(kind,'')):
            raise ValueError('Configure a pasta do estado ou informe a pasta deste município.')
        if not root.is_dir(): raise ValueError('Pasta não encontrada: '+str(root))
        if override:
            target=root
        else:
            matches=[p for p in root.iterdir() if p.is_dir() and norm(p.name)==norm(city)]
            if norm(root.name)==norm(city): matches=[root]
            if len(matches)!=1: raise ValueError('Pasta do município não encontrada ou ambígua. Informe o caminho manualmente.')
            target=matches[0]
        if (target/'2004').is_dir(): target=target/'2004'
        return target.resolve()

    def photo_upload(self,key,kind,relative,content,batch):
        with self.connect() as db:
            if not db.execute('SELECT 1 FROM municipalities WHERE key=?',(key,)).fetchone():raise ValueError('Município não encontrado.')
        if kind not in ('original','restored'):raise ValueError('Tipo de fotos inválido.')
        if not re.fullmatch(r'[a-f0-9]{24}',batch):raise ValueError('Lote inválido.')
        parts=relative.replace('\\','/').split('/')
        if any(x in ('','..','.') or ':' in x for x in parts):raise ValueError('Nome de arquivo inválido.')
        if Path(relative).suffix.lower() not in IMAGE_EXT:raise ValueError('Use arquivos de imagem.')
        try:fitz.Pixmap(content)
        except Exception:raise ValueError('Imagem inválida: '+parts[-1])
        folder=self.folder/'fotos_importadas'/norm(key)/kind/batch
        target=folder.joinpath(*parts);target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(content)
        return {'ok':True}

    def finish_photo_upload(self,key,kind,batch):
        if kind not in ('original','restored') or not re.fullmatch(r'[a-f0-9]{24}',batch):raise ValueError('Lote inválido.')
        folder=self.folder/'fotos_importadas'/norm(key)/kind/batch
        return self.load_originals(key,str(folder)) if kind=='original' else self.load_restored(key,str(folder))

    def pick_folder(self,initial=''):
        initial=str(initial or '')
        if not Path(initial).is_dir():initial=str(ROOT)
        script="""import json,sys,tkinter as tk
from tkinter import filedialog
root=tk.Tk();root.withdraw();root.attributes('-topmost',True)
try:
    folder=filedialog.askdirectory(parent=root,title='JP Minato — selecione a pasta de fotos',initialdir=sys.argv[1],mustexist=True)
    print(json.dumps({'path':folder,'cancelled':not bool(folder)}))
finally:root.destroy()
"""
        try:
            options={'creationflags':subprocess.CREATE_NO_WINDOW} if sys.platform=='win32' else {}
            result=subprocess.run([sys.executable,'-c',script,initial],capture_output=True,text=True,encoding='utf-8',timeout=300,**options)
            if result.returncode:raise ValueError('Não foi possível abrir o seletor de pastas. Cole o caminho completo no campo; a importação começa ao sair do campo.')
            choice=json.loads(result.stdout)
        except subprocess.TimeoutExpired:raise ValueError('Seleção de pasta encerrada por tempo limite. Tente novamente.')
        if not choice.get('cancelled') and not Path(choice['path']).is_dir():raise ValueError('Pasta selecionada não encontrada.')
        return choice

    def scan_images(self,folder):
        results=[]
        for p in sorted(folder.rglob('*')):
            if not p.is_file() or p.suffix.lower() not in IMAGE_EXT: continue
            relative=p.relative_to(folder)
            if any(re.fullmatch(r'20\d{2}',part) and part!='2004' for part in relative.parts[:-1]): continue
            results.append({'path':str(p.resolve()),'name':p.stem.replace('_',' '),'match':norm(p.stem),'relative':str(relative),'identity':'/'.join([norm(x) for x in relative.parts[:-1]]+[norm(p.stem)])})
        if not results: raise ValueError('Nenhuma imagem de 2004 encontrada nesta pasta.')
        return results

    def load_originals(self,key,override=None):
        source=self.photo_folder(key,override)
        folder,files=data_storage.folder_import(self,key,source,'original')
        project=self.project(key)
        old={p['original']:p for p in project.get('pairs',[])}
        identities={}
        for pair in project.get('pairs',[]):identities.setdefault(pair['identity'],[]).append(pair)
        new=[]
        for file in files:
            matches=identities.get(file['identity'],[])
            previous=old.get(file['path']) or (matches[0] if len(matches)==1 else {})
            new.append({**file,'id':previous.get('id',secrets.token_hex(8)),'original':file['path'],'restored':previous.get('restored'),'selected':previous.get('selected',True),'approved':False,'name':previous.get('name',file['name'])})
        project['pairs']=new
        project.setdefault('import_sources',{})['original']=str(source)
        project['original_folder']=str(folder);project['dirty']=True
        self.save_project(key,project)
        return self.public_project(key)

    def load_restored(self,key,override=None):
        source=self.photo_folder(key,override,kind='restored')
        project=self.project(key)
        if not project['pairs']: raise ValueError('Carregue as fotos originais primeiro.')
        folder,files=data_storage.folder_import(self,key,source,'restored')
        identities={}
        by_name={}
        original_names={}
        for p in project['pairs']: original_names[p['match']]=original_names.get(p['match'],0)+1
        for f in files:
            identities.setdefault(f['identity'],[]).append(f)
            by_name.setdefault(f['match'],[]).append(f)
        used=set()
        for pair in project['pairs']:
            candidates=identities.get(pair['identity'],[])
            if not candidates and original_names[pair['match']]==1: candidates=by_name.get(pair['match'],[])
            match=candidates[0] if len(candidates)==1 and candidates[0]['path'] not in used else None
            pair['restored']=match['path'] if match else None
            pair['approved']=False
            pair['problem']='' if match else ('Nome duplicado / associação manual' if candidates or original_names[pair['match']]>1 else 'Tratada não encontrada')
            if match:used.add(match['path'])
        project.setdefault('import_sources',{})['restored']=str(source)
        project['restored_files']=files;project['restored_folder']=str(folder);project['dirty']=True
        self.save_project(key,project)
        return self.public_project(key)

    def token_for(self,path):
        token=secrets.token_hex(16);self.files[token]=str(path);return token

    def public_project(self,key):
        p=self.project(key)
        for pair in p.get('pairs',[]):
            pair['original_url']='/file/'+self.token_for(pair['original'])
            pair['restored_url']='/file/'+self.token_for(pair['restored']) if pair.get('restored') else None
        if p.get('pdf'):p['pdf_url']='/file/'+self.token_for(p['pdf'])
        if p.get('crest'):p['crest_url']='/file/'+self.token_for(p['crest'])
        return p

    def update_pairs(self,key,payload):
        p=self.project(key)
        before=json.dumps(p,sort_keys=True,ensure_ascii=False)
        actual={x['id']:x for x in p['pairs']}
        permitted={x['path'] for x in p.get('restored_files',[])}
        for update in payload.get('pairs',[]):
            if update['id'] not in actual:raise ValueError('Foto não encontrada.')
            pair=actual[update['id']]
            restored=update.get('restored') or None
            if restored and restored not in permitted and restored!=pair.get('restored'):raise ValueError('Foto tratada fora da pasta importada.')
            pair.update({k:update[k] for k in ('name','selected','approved') if k in update})
            pair['restored']=restored
            if not restored:pair['approved']=False
        selected=[x.get('restored') for x in p['pairs'] if x.get('selected',True) and x.get('restored')]
        if len(selected)!=len(set(selected)):raise ValueError('Uma foto tratada foi associada a mais de uma pessoa.')
        p['period']=str(payload.get('period',p.get('period','2005–2008')))
        if json.dumps(p,sort_keys=True,ensure_ascii=False)!=before:p['dirty']=True
        self.save_project(key,p)

    def crest(self,key,content,name):
        extension=Path(name).suffix.lower()
        if extension not in {'.png','.jpg','.jpeg','.webp'}:raise ValueError('Use um brasão em PNG, JPG ou WebP.')
        content=crest_capture.as_png(content);extension='.png'
        folder=self.folder/'brasoes';folder.mkdir(exist_ok=True)
        target=folder/(norm(key)+'_'+secrets.token_hex(6)+extension)
        with self.lock:
            target.write_bytes(content);p=self.project(key)
            if p.get('crest'):p.setdefault('crest_versions',[]).append(p['crest'])
            p['crest']=str(target);p['dirty']=True;self.save_project(key,p)

    def export_originals(self,key):
        item=self.item(key);p=self.project(key)
        pairs=[x for x in p['pairs'] if x.get('selected',True)]
        if not pairs:raise ValueError('Selecione ao menos uma foto original.')
        folder=self.folder/'zips';folder.mkdir(exist_ok=True)
        target=folder/(safe_filename(item['fields']['Município']+'_'+item['fields']['UF']+'_2004_originais')+'.zip')
        with zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED) as archive:
            for pair in pairs:
                file=Path(pair['original'])
                if not file.is_file():raise ValueError('Foto movida ou apagada: '+str(file))
                archive.write(file,Path(pair['relative']).as_posix())
            archive.writestr('manifest.json',json.dumps({'municipal_code':item['code'],'municipality_key':key,'municipio':item['fields']['Município'],'uf':item['fields']['UF'],'ano':2004,'fotos':[{'id':x['id'],'nome':x['name'],'arquivo':Path(x['relative']).as_posix()} for x in pairs]},ensure_ascii=False,indent=2))
        return {'url':'/file/'+self.token_for(target),'filename':target.name}

    def export_contacts(self,keys):
        wanted=set(keys)
        items=[x for x in self.items() if x['key'] in wanted]
        if not items:raise ValueError('Selecione municípios ou ajuste os filtros antes de exportar.')
        columns=list(dict.fromkeys(k for item in items for k in item['fields']))
        extras=['Setor comercial','Situação comercial','DDD municipal','Etapa do material','Último contato','Observações comerciais']
        workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Municípios'
        sheet.append(columns+extras)
        for item in items:
            local=item['local']
            values=[item['fields'].get(k,'') for k in columns]+[local.get('sector',''),local.get('commercial','Não contatado'),item['ddd'],item['stage'],local.get('last_contact',''),local.get('notes','')]
            sheet.append([str(v) if v is not None else '' for v in values])
            # Export external text as literal values, including strings beginning with '='.
            for cell in sheet[sheet.max_row]:cell.data_type='s'
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=30
        folder=self.folder/'exportacoes';folder.mkdir(exist_ok=True)
        target=folder/f'Contatos_JP_Minato_{datetime.now():%Y%m%d_%H%M%S}_{secrets.token_hex(2)}.xlsx'
        workbook.save(target)
        return {'url':'/file/'+self.token_for(target),'filename':target.name,'count':len(items)}

    def catalogue_info(self):
        saved=self.setting('catalogue',{})
        path=catalogue125.standard_path(self)
        filename=saved.get('filename') if saved.get('layout_version') in ('1.33','1.33.1') else 'Catálogo base - páginas 3 a 10.pdf'
        if not path.is_file():return {'filename':filename,'available':False,'pages':0,'url':None,'warning':'PDF padrão não encontrado. Seus cadastros continuam disponíveis. Importe novamente o PDF nas configurações.'}
        try:
            with fitz.open(path) as doc:
                if doc.needs_pass:raise ValueError('PDF com senha.')
                pages=len(doc)
        except Exception:return {'filename':filename,'available':False,'pages':0,'url':None,'warning':'Não foi possível ler o PDF padrão. Importe novamente o arquivo nas configurações.'}
        return {'filename':filename,'available':True,'pages':pages,'url':'/file/'+self.token_for(path)}

    def upload_catalogue(self,content,name):
        if Path(name).suffix.lower()!='.pdf':raise ValueError('Escolha um arquivo PDF.')
        try:
            with fitz.open(stream=content,filetype='pdf') as doc:
                if doc.needs_pass:raise ValueError('Use um PDF sem senha.')
                if len(doc)<1:raise ValueError('O PDF precisa conter pelo menos uma página.')
                doc[0].get_pixmap(matrix=fitz.Matrix(.1,.1))
        except Exception as error:raise ValueError('Não foi possível ler o PDF padrão: '+str(error))
        folder=self.folder/'modelos';folder.mkdir(exist_ok=True)
        target=folder/('restante_'+secrets.token_hex(8)+'.pdf');target.write_bytes(content)
        self.set_setting('catalogue',{'path':str(target),'filename':Path(name).name,'layout_version':'1.33.1'})
        with self.connect() as db:keys=[row[0] for row in db.execute('SELECT key FROM projects')]
        for key in keys:self.invalidate(key)
        return self.catalogue_info()

    def personalized_cover(self,city,uf,crest,draft):
        return catalogue125.cover(self,city,uf,crest,draft)

    def pdf(self,key,draft=False):
        item=self.item(key);p=self.project(key)
        pairs=[x for x in p['pairs'] if x.get('selected',True)]
        if not pairs:raise ValueError('Selecione ao menos uma foto.')
        if not draft:
            if not p.get('crest'):raise ValueError('Cadastre o brasão antes de gerar o PDF final.')
            if any(not x.get('restored') for x in pairs):raise ValueError('Importe as fotos tratadas correspondentes a todas as pessoas selecionadas antes do PDF final.')
        for pair in pairs:
            for kind in ('original','restored'):
                if pair.get(kind) and not Path(pair[kind]).is_file():raise ValueError('Foto movida ou apagada: '+pair[kind])
        city=item['fields']['Município'];uf=item['fields']['UF']
        doc=catalogue125.build(self,item,p,pairs,draft)
        personalized=len(doc)-1
        standard_path=catalogue125.standard_path(self)
        if not standard_path.is_file():raise ValueError('Importe o PDF padrão novamente nas configurações.')
        with fitz.open(standard_path) as standard:
            standard_pages=len(standard);doc.insert_pdf(standard)
        # Preserve the exact page size supplied in Catalogo Base.pdf.
        total_pages=len(doc)
        folder=self.folder/'pdfs'/secrets.token_hex(12);folder.mkdir(parents=True,exist_ok=True)
        filename=safe_filename(f"{'PREVIA - ' if draft else ''}{city} Vereadores 2005 a 2008.pdf")
        target=folder/filename;doc.save(target,garbage=4,deflate=True);doc.close()
        if not draft:p['pdf']=str(target);p['dirty']=False;p['generated_at']=now();p['catalogue_version']=1;p['layout_version']='1.33.1';self.save_project(key,p)
        return {'url':'/file/'+self.token_for(target),'filename':filename,'pages':total_pages,'cover_pages':1,'personalized_pages':personalized,'standard_pages':standard_pages}

STORE=None
TOKEN=secrets.token_urlsafe(32)
CAPTURE_PUBLIC=remote_access.CAPTURE_ROUTES | whatsapp_sender129.WORKER_ROUTES | {'/whatsapp-escritorio','/capturar','/api/login','/api/login-options','/api/gallery-capture','/api/health','/api/server-stop'}
ACCESS=None
REMOTE_ONLINE=False

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def response(self,status,value,ctype='application/json; charset=utf-8'):
        if self.command=='POST' and 200<=status<300 and operators131.actor() and STORE and urlparse(self.path).path not in {'/api/smtp-test','/api/inbox-sync','/api/restoration-preview','/api/capture-preview'}:
            for key in getattr(self,'audit_keys',[]) or [None]:operators131.audit(STORE,urlparse(self.path).path,key)
        content=json.dumps(value,ensure_ascii=False).encode() if isinstance(value,(dict,list)) else value
        self.send_response(status);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(content)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('Referrer-Policy','same-origin');self.send_header('X-Frame-Options','DENY');self.end_headers()
        try:self.wfile.write(content)
        except (ConnectionAbortedError,ConnectionResetError,BrokenPipeError):pass
    def gate(self,path):
        if ACCESS is None:return True
        if not ACCESS.valid_host(self.headers.get('Host',''),self.server.server_port):
            self.response(403,{'error':'Endereço do servidor não autorizado.'});return False
        if path in CAPTURE_PUBLIC or path.startswith('/static/') or ACCESS.authenticated(self.headers.get('Cookie')) and ACCESS.operator(self.headers.get('Cookie')):return True
        if path=='/':self.response(200,(ROOT/'web'/'login.html').read_bytes(),'text/html; charset=utf-8')
        else:self.response(401,{'error':'Sessão expirada. Entre novamente no JP Minato.'})
        return False
    def do_GET(self):
        try:
            route=urlparse(self.path);query=parse_qs(route.query)
            if not self.gate(route.path):return
            if self.retired(route.path):return
            if route.path=='/api/login-options':return self.response(200,{'password_required':bool(ACCESS and ACCESS.config.get('enabled'))})
            if route.path=='/api/dashboard':return self.response(200,dashboard132.summary(STORE,query,ACCESS.operator(self.headers.get('Cookie')) if ACCESS else None))
            if route.path=='/api/inbox-config':return self.response(200,inbox131.public(STORE))
            if route.path=='/api/inbox':return self.response(200,inbox131.listing(STORE,query))
            if route.path=='/api/inbox-message':return self.response(200,inbox131.detail(STORE,query['id'][0]))
            if route.path=='/api/operator-history':return self.response(200,operators131.history(STORE,query.get('key',[None])[0]))
            if route.path=='/api/inbox-attachment':
                raw,name=inbox131.attachment(STORE,query['id'][0],query['index'][0]);return self.attachment(raw,name)
            if route.path=='/api/health':return self.response(200,{'app':'JP Minato','version':'1.36','capabilities':['operator_sessions_v1','imap_inbox_v1'],'instance':ACCESS.instance if ACCESS else ''})
            if route.path=='/api/access':return self.response(200,ACCESS.public(self.client_address[0]!='127.0.0.1',REMOTE_ONLINE) if ACCESS else {'enabled':False,'remote':False})
            if route.path=='/whatsapp-escritorio':return self.response(200,(ROOT/'web'/'wa-office-start.html').read_bytes(),'text/html; charset=utf-8')
            if route.path=='/api/wa-office':return self.response(200,whatsapp_sender129.public(STORE))
            if route.path=='/api/wa-office-status':return self.response(200,whatsapp_sender129.lookup(STORE,query['id'][0]))
            if route.path=='/capturar':return self.response(200,(ROOT/'web'/'capture-start.html').read_bytes(),'text/html; charset=utf-8')
            if route.path=='/':
                page=(ROOT/'web'/'index.html').read_text('utf-8').replace('__TOKEN__',TOKEN)
                return self.response(200,page.encode(),'text/html; charset=utf-8')
            if route.path.startswith('/file/'):
                path=STORE.files.get(route.path.split('/')[-1])
                if not path or not Path(path).is_file():return self.response(404,{'error':'Arquivo não disponível. Recarregue as fotos.'})
                if Path(path).suffix.lower() in {'.tif','.tiff','.bmp'}:
                    return self.response(200,fitz.Pixmap(path).tobytes('png'),'image/png')
                return self.response(200,Path(path).read_bytes(),mimetypes.guess_type(path)[0] or 'application/octet-stream')
            if route.path=='/api/extractor':return self.response(200,contact_extractor.public_config(STORE))
            if route.path=='/api/extractor-job':return self.response(200,contact_extractor.job(STORE,query['job'][0]))
            if route.path=='/api/extractor-prompt':return self.response(200,contact_extractor.extractor_prompt.snapshot(STORE))
            if route.path=='/api/extractor-rules':return self.response(200,contact_extractor.rules(STORE))
            if route.path=='/api/capture-changes':
                after=int(query.get('after',['0'])[0])
                wait=max(0,min(20,float(query.get('wait',['0'])[0])))
                with STORE.capture_condition:
                    if STORE.capture_sequence==after and wait:STORE.capture_condition.wait_for(lambda:STORE.capture_sequence!=after,timeout=wait)
                    sequence=STORE.capture_sequence
                    keys=list(dict.fromkeys(k for n,k in STORE.capture_events if n>after))
                    updates=[x.copy() for x in STORE.capture_updates if x['sequence']>after]
                return self.response(200,{'sequence':sequence,'keys':keys,'updates':updates,'items':[STORE.item(k) for k in keys],'projects':{k:STORE.public_project(k) for k in keys if any(u['key']==k and u['field']=='Brasão' for u in updates)}})
            if route.path=='/api/smtp-config':return self.response(200,smtp_mail127.public(STORE))
            if route.path=='/api/smtp-status':return self.response(200,smtp_mail127.public_attempt(smtp_mail127.lookup(STORE,query['id'][0])))
            if route.path=='/api/data':return self.response(200,{'items':STORE.items(),'routes':STORE.routes(),'roots':STORE.root_settings(),'sectors':STORE.setting('sectors',[]),'tse_states':STORE.tse_public(),'catalogue':STORE.catalogue_info(),'tse_national':STORE.setting('tse_national'),'access':ACCESS.public(self.client_address[0]!='127.0.0.1',REMOTE_ONLINE) if ACCESS else {'enabled':False,'remote':False},'operator':ACCESS.operator(self.headers.get('Cookie')) if ACCESS else None,'storage':{'database':str(STORE.db_path.resolve()),'folder':str(STORE.folder.resolve())}})
            if route.path.startswith('/static/'):
                relative=route.path[len('/static/'):]
                base=(ROOT/'web').resolve();path=(base/relative).resolve()
                if not path.is_relative_to(base) or not path.is_file():return self.response(404,{'error':'Recurso não encontrado.'})
                return self.response(200,path.read_bytes(),mimetypes.guess_type(path)[0] or 'application/octet-stream')
            if route.path=='/api/gallery-assessment':return self.response(200,galleries135.public(STORE,query['key'][0]))
            if route.path=='/api/visit-routes':return self.response(200,{'routes':galleries135.routes(STORE)})
            if route.path=='/api/routes':return self.response(200,{'routes':STORE.routes()})
            if route.path=='/api/message-templates':return self.response(200,{'templates':communication.templates(STORE)})
            if route.path=='/api/contact-openings':return self.response(200,communication.opened(STORE,query['key'][0]))
            if route.path=='/api/map':
                path=ROOT/'assets'/'maps'/'RS.json'
                return self.response(200,path.read_bytes(),'application/json; charset=utf-8')
            if route.path=='/api/phone-check':return self.response(200,STORE.check_phone(query['key'][0],query['number'][0]))
            if route.path=='/api/storage':return self.response(200,data_storage.status(STORE))
            if route.path=='/api/browser-verifier':return self.response(200,whatsapp_browser.public(STORE))
            if route.path=='/api/google-contacts':return self.response(200,google_contacts.public(STORE))
            if route.path=='/api/history':return self.response(200,STORE.interaction_history(query['key'][0]))
            if route.path=='/api/media-job':return self.response(200,STORE.import_status(query['job'][0]))
            if route.path=='/api/restoration':return self.response(200,restoration_ai.public(STORE))
            if route.path=='/api/whatsapp-verification':return self.response(200,whatsapp_verify.public(STORE))
            if route.path=='/api/tse-job':return self.response(200,STORE.import_status(query['job'][0]))
            if route.path=='/api/project':return self.response(200,STORE.public_project(query['key'][0]))
            return self.response(404,{'error':'Rota não encontrada.'})
        except Exception as e:self.response(400,{'error':str(e)})
    def attachment(self,raw,name):
        from urllib.parse import quote
        self.send_response(200);self.send_header('Content-Type','application/octet-stream');self.send_header('Content-Length',str(len(raw)));self.send_header('X-Content-Type-Options','nosniff');self.send_header('Content-Disposition',"attachment; filename*=UTF-8''"+quote(name,safe=''));self.end_headers();self.wfile.write(raw)
    def retired(self,path):
        if path.startswith(('/api/extractor','/api/browser-verifier','/api/verifier-')) or path in {'/api/whatsapp-verification','/api/whatsapp-settings','/api/whatsapp-test','/api/whatsapp-enqueue','/api/whatsapp-reset'}:
            self.response(410,{'error':'Funcionalidade retirada na versão 1.32. Dados anteriores preservados.'});return True
        return False
    def do_POST(self):
        who=ACCESS.operator(self.headers.get('Cookie')) if ACCESS else None
        binding=STORE.capture_sessions.get(self.headers.get('X-Capture-Token',''),{}) if STORE else {}
        if binding.get('operator'):who=binding['operator']
        token=operators131.CURRENT.set(who)
        self.capture_key131=binding.get('key')
        self.audit_keys=[]
        try:self.post_impl131()
        finally:operators131.CURRENT.reset(token)
    def post_impl131(self):
        try:
            access_path=urlparse(self.path).path
            if not self.gate(access_path):return
            if self.retired(access_path):return
            if access_path=='/api/login':
                if self.headers.get('Origin')!='http://'+self.headers.get('Host',''):return self.response(403,{'error':'Origem inválida.'})
                length=int(self.headers.get('Content-Length','0'))
                if not 0<length<=2048:raise ValueError('Solicitação inválida.')
                login=json.loads(self.rfile.read(length))
                if login.get('operator') not in operators131.NAMES:raise ValueError('Selecione Jeff, Lara ou Pai para entrar.')
                token=ACCESS.login(login.get('password',''),self.client_address[0],login['operator'])
                self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Cache-Control','no-store');self.send_header('Set-Cookie','jp_session='+token+'; HttpOnly; SameSite=Strict; Path=/; Max-Age=43200');self.end_headers();self.wfile.write(b'{"ok":true}');return
            if access_path=='/api/server-stop':
                if self.client_address[0]!='127.0.0.1' or not ACCESS or not secrets.compare_digest(self.headers.get('X-Server-Control',''),ACCESS.control):return self.response(403,{'error':'Controle local inválido.'})
                self.response(200,{'ok':True});threading.Thread(target=self.server.shutdown,daemon=True).start();return
            capture_route=urlparse(self.path).path in whatsapp_sender129.WORKER_ROUTES or urlparse(self.path).path in ('/api/verifier-next','/api/verifier-result','/api/capture-connect','/api/capture-navigate','/api/capture-preview','/api/capture-save','/api/crest-save','/api/gallery-capture')
            if capture_route:
                origin=self.headers.get('Origin','');extension_id=self.headers.get('X-JP-Extension','')
                trusted_origin=bool(re.fullmatch(r'chrome-extension://[a-p]{32}',origin))
                trusted_extension=not origin and bool(re.fullmatch(r'[a-p]{32}',extension_id))
                if not (trusted_origin or trusted_extension):return self.response(403,{'error':'A captura deve ser iniciada pela extensão JP Minato atualizada. Recarregue a extensão no Chrome.'})
                if trusted_origin and extension_id and origin!='chrome-extension://'+extension_id:return self.response(403,{'error':'Identificação da extensão inválida.'})
            elif self.headers.get('X-App-Token')!=TOKEN:return self.response(403,{'error':'Reabra o aplicativo para continuar.'})
            length=int(self.headers.get('Content-Length','0'))
            if length<0 or self.headers.get('Transfer-Encoding'):raise ValueError('Tamanho de solicitação inválido.')
            route=urlparse(self.path);query=parse_qs(route.query)
            limit=250 if route.path in ('/api/tse-sheet','/api/tse-start') else 30
            if length>limit*1024*1024:raise ValueError(f'Arquivo maior que {limit} MB.')
            self.audit_keys=[query['key'][0]] if query.get('key') else []
            raw=self.rfile.read(length)
            if route.path=='/api/media-upload':return self.response(200,municipal_media.upload(STORE,query['batch'][0],query['name'][0],raw))
            if route.path=='/api/photo-upload':return self.response(200,STORE.photo_upload(query['key'][0],query['kind'][0],query['name'][0],raw,query['batch'][0]))
            if route.path=='/api/catalogue-upload':return self.response(200,STORE.upload_catalogue(raw,query.get('name',['restante.pdf'])[0]))
            if route.path=='/api/gallery-photo':return self.response(200,galleries135.photo(STORE,query['key'][0],raw))
            if route.path=='/api/import':return self.response(200,STORE.import_xlsx(raw))
            if route.path=='/api/tse-start':return self.response(200,STORE.start_tse_import(query['uf'][0],raw,query['name'][0]))
            if route.path=='/api/tse-sheet':return self.response(200,STORE.tse_import(query['uf'][0],raw,query['name'][0]))
            if route.path=='/api/tse-photo':return self.response(200,STORE.tse_photo_upload(query['uf'][0],query['name'][0],raw,query['batch'][0]))
            if route.path=='/api/crest':
                STORE.crest(query['key'][0],raw,query.get('name',['brasao.png'])[0]);return self.response(200,{'ok':True})
            body=json.loads(raw or b'{}');key=body.get('key')
            self.audit_keys=[key] if key else body.get('keys',[]) if isinstance(body.get('keys'),list) else [self.capture_key131] if self.capture_key131 else []
            if self.client_address[0]!='127.0.0.1' and route.path in ('/api/pick-folder','/api/open-folder'):raise ValueError('Escolha uma pasta pelo navegador para enviar os arquivos desta máquina.')
            if route.path=='/api/logout':
                if ACCESS:ACCESS.logout(self.headers.get('Cookie'))
                return self.response(200,{'ok':True})
            if route.path=='/api/desk-task':result=dashboard132.save_task(STORE,body)
            elif route.path=='/api/desk-case':result=dashboard132.save_case(STORE,body)
            elif route.path=='/api/desk-call':result=dashboard132.call(STORE,body)
            elif route.path=='/api/route-owner':result=operators131.assign(STORE,body)
            elif route.path=='/api/inbox-config':result=inbox131.configure(STORE,body)
            elif route.path=='/api/inbox-sync':result=municipal_media.background(STORE,lambda progress:inbox131.sync(STORE,body.get('test') is True))
            elif route.path=='/api/inbox-change':result=inbox131.change(STORE,body)
            elif route.path=='/api/inbox-reply':result=inbox131.reply(STORE,body)
            elif route.path=='/api/wa-office-config':result=whatsapp_sender129.configure(STORE,body)
            elif route.path=='/api/wa-office-pair':result=whatsapp_sender129.pair(STORE,self.client_address[0])
            elif route.path=='/api/wa-office-connect':result=whatsapp_sender129.connect(STORE,body,self.client_address[0],self.headers.get('X-JP-Extension','') or self.headers.get('Origin','').removeprefix('chrome-extension://'))
            elif route.path=='/api/wa-office-worker':result=whatsapp_sender129.worker(STORE,self.headers.get('X-Capture-Token'),body,self.client_address[0],self.headers.get('X-JP-Extension','') or self.headers.get('Origin','').removeprefix('chrome-extension://'))
            elif route.path=='/api/wa-office-send':result=whatsapp_sender129.prepare(STORE,key,body)
            elif route.path=='/api/wa-office-cancel':result=whatsapp_sender129.cancel(STORE,body.get('id'))
            elif route.path=='/api/wa-office-block':result=whatsapp_sender129.change_block(STORE,body)
            elif route.path=='/api/smtp-config':result=smtp_mail127.configure(STORE,body)
            elif route.path=='/api/smtp-test':result=smtp_mail127.test_connection(STORE)
            elif route.path=='/api/smtp-send':result=smtp_mail127.prepare(STORE,key,body)
            elif route.path=='/api/media-export':result=municipal_media.background(STORE,lambda progress:municipal_media.export_batch(STORE,body.get('keys',[]),progress))
            elif route.path=='/api/media-import':
                path=body.get('path') or str(STORE.folder/'lotes_municipais'/str(body.get('batch','')))
                if not body.get('path') and not re.fullmatch(r'[a-f0-9]{24}',str(body.get('batch',''))):raise ValueError('Lote inválido.')
                result=municipal_media.background(STORE,lambda progress:municipal_media.import_folders(STORE,path,body.get('kind','original'),body.get('replace') is True,progress))
            elif route.path=='/api/restoration-settings':result=restoration_ai.save_config(STORE,body)
            elif route.path=='/api/restoration-test':result=restoration_ai.test_connection(STORE)
            elif route.path=='/api/restoration-preview':
                result=restoration_ai.preview(STORE,body);result.pop('tasks')
            elif route.path=='/api/restoration-start':result=restoration_ai.enqueue(STORE,body)
            elif route.path=='/api/restoration-control':result=restoration_ai.control(STORE,body)
            elif route.path=='/api/whatsapp-settings':result=whatsapp_verify.save_config(STORE,body)
            elif route.path=='/api/whatsapp-test':result=whatsapp_verify.test_connection(STORE)
            elif route.path=='/api/whatsapp-enqueue':result=whatsapp_verify.enqueue(STORE,body.get('keys',[]),manual=True)
            elif route.path=='/api/whatsapp-reset':result=whatsapp_verify.reset(STORE,body.get('number'))
            elif route.path=='/api/extractor-prompt':result=contact_extractor.extractor_prompt.save(STORE,body)
            elif route.path=='/api/extractor-settings':result=contact_extractor.save_settings(STORE,body)
            elif route.path=='/api/extractor-resume':result=contact_extractor.resume(STORE,body.get('job'))
            elif route.path=='/api/extractor-start':result=contact_extractor.start(STORE,body.get('keys',[]),body.get('mode'),body.get('label',''))
            elif route.path=='/api/extractor-cancel':result=contact_extractor.cancel(STORE,body.get('job'))
            elif route.path=='/api/extractor-test':result=STORE.contact_engine.call('/api/health')
            elif route.path=='/api/extractor-rules':result=contact_extractor.save_rules(STORE,body)
            elif route.path=='/api/extractor-accept':result=contact_extractor.accept_candidate(STORE,key,body.get('candidate_id'))
            elif route.path=='/api/capture-connect':result=STORE.connect_capture(body.get('ticket'),body.get('tab_id'))
            elif route.path=='/api/capture-navigate':result=capture_routes.advance(STORE,self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/capture-preview':result=STORE.capture_preview(self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/capture-save':result=STORE.save_capture(self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/crest-start':result=crest_capture.start(STORE,key)
            elif route.path=='/api/crest-save':result=crest_capture.capture(STORE,self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/capture-start':result=STORE.start_capture(key,body.get('url'))
            elif route.path=='/api/message-templates':result=communication.save_templates(STORE,body)
            elif route.path=='/api/contact-opened':result=communication.opened(STORE,key,body)
            elif route.path=='/api/settings':
                STORE.set_setting('roots',body['roots']);STORE.set_setting('sectors',body['sectors']);result={'ok':True}
            elif route.path=='/api/photo-finish':result=STORE.finish_photo_upload(key,body['kind'],body['batch'])
            elif route.path=='/api/pick-folder':result=STORE.pick_folder(body.get('initial',''))
            elif route.path=='/api/tse-process':result=STORE.tse_process(body['uf'],body.get('path'))
            elif route.path=='/api/storage-collect':result=municipal_media.background(STORE,lambda progress:data_storage.migrate(STORE,progress))
            elif route.path=='/api/browser-verifier-configure':result=whatsapp_browser.configure(STORE,body)
            elif route.path=='/api/browser-verifier-enqueue':result=whatsapp_browser.enqueue(STORE,body.get('keys'))
            elif route.path=='/api/browser-verifier-start':result=whatsapp_browser.start(STORE)
            elif route.path=='/api/browser-verifier-control':result=whatsapp_browser.control(STORE,body)
            elif route.path=='/api/verifier-next':result=whatsapp_browser.next_task(STORE,self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/verifier-result':result=whatsapp_browser.result(STORE,self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/google-configure':result=google_contacts.configure(STORE,body)
            elif route.path=='/api/google-authorize':
                if self.client_address[0]!='127.0.0.1':raise ValueError('Autorize a conta no navegador do computador do escritório. Depois, use a sincronização de qualquer máquina pela VPN.')
                result=google_contacts.authorize(STORE)
            elif route.path=='/api/google-csv':
                raw=google_contacts.csv_export(STORE,body.get('keys'));path=STORE.folder/'exportacoes'/('Google-Contatos-'+secrets.token_hex(6)+'.csv');path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(raw);result={'url':'/file/'+STORE.token_for(path),'filename':path.name}
            elif route.path=='/api/google-preview':result=municipal_media.background(STORE,lambda progress:google_contacts.prepare(STORE,body.get('keys'),progress))
            elif route.path=='/api/google-sync':result=municipal_media.background(STORE,lambda progress:google_contacts.apply(STORE,body,progress))
            elif route.path=='/api/google-new-president':result=google_contacts.new_president(STORE,key)
            elif route.path=='/api/whatsapp-send-verify':result=whatsapp_send126.verify(STORE,key,body)
            elif route.path=='/api/interaction-correct':result=STORE.correct_interaction(key,body)
            elif route.path=='/api/interaction':result=STORE.record_interaction(key,body)
            elif route.path=='/api/gallery-assessment':result=galleries135.save(STORE,body)
            elif route.path=='/api/gallery-link':result=galleries135.link(STORE,body)
            elif route.path=='/api/gallery-remove':result=galleries135.remove(STORE,body)
            elif route.path=='/api/gallery-capture':result=galleries135.capture(STORE,self.headers.get('X-Capture-Token'),body)
            elif route.path=='/api/visit-route-save':result=galleries135.visit_save(STORE,body)
            elif route.path=='/api/visit-route-remove':result=galleries135.visit_remove(STORE,body)
            elif route.path=='/api/route-save':result=STORE.save_route(body)
            elif route.path=='/api/route-delete':result=STORE.delete_route(body)
            elif route.path=='/api/phone':result=STORE.add_phone(key,body)
            elif route.path=='/api/review':result=editing_activity.review(STORE,key)
            elif route.path=='/api/save':
                STORE.save_item(key,body.get('fields',{}),body.get('local',{}),manual=True)
                result={'ok':True,'item':STORE.item(key)}
            elif route.path=='/api/originals':result=STORE.load_originals(key,body.get('path'))
            elif route.path=='/api/restored':result=STORE.load_restored(key,body.get('path'))
            elif route.path=='/api/pairs':
                with STORE.lock:STORE.update_pairs(key,body)
                result={'ok':True}
            elif route.path=='/api/pdf':result=STORE.pdf(key,body.get('draft',False))
            elif route.path=='/api/export-contacts':result=STORE.export_contacts(body.get('keys',[]))
            elif route.path=='/api/export-originals':result=STORE.export_originals(key)
            elif route.path=='/api/open-folder':
                p=STORE.project(key);folder=p.get('original_folder')
                if not folder or not Path(folder).is_dir():raise ValueError('Carregue a pasta de originais primeiro.')
                if sys.platform=='win32':os.startfile(folder)
                else:webbrowser.open(Path(folder).as_uri())
                result={'ok':True}
            elif route.path=='/api/batch-sector':
                sector=body['sector']
                if sector not in STORE.setting('sectors',[]):raise ValueError('Cadastre o setor nas configurações.')
                for key in body['keys']:STORE.save_item(key,{}, {'sector':sector},manual=True)
                result={'ok':True}
            else:return self.response(404,{'error':'Rota não encontrada.'})
            self.response(200,result)
        except Exception as e:self.response(400,{'error':str(e)})

def main():
    global STORE, ACCESS, REMOTE_ONLINE
    config=remote_access.load_config(DATA)
    ACCESS=remote_access.Access(config)
    server=ThreadingHTTPServer(('127.0.0.1',8765),Handler)
    server.daemon_threads=True
    runtime=DATA/'servidor-em-execucao.json'
    STORE=Store()
    restoration_ai.start_worker(STORE);inbox131.start_worker(STORE)
    if not STORE.items() and (ROOT/'assets'/'contatos.xlsx').exists():STORE.import_xlsx((ROOT/'assets'/'contatos.xlsx').read_bytes())
    remote_access.write_json(runtime,{'pid':os.getpid(),'instance':ACCESS.instance,'control':ACCESS.control})
    stop=threading.Event();remote_servers=[]
    def vpn_listener():
        global REMOTE_ONLINE
        while not stop.is_set():
            try:
                remote=ThreadingHTTPServer((config['vpn_ip'],8765),Handler);remote.daemon_threads=True
                if stop.is_set():remote.server_close();break
                remote_servers.append(remote);REMOTE_ONLINE=True
                remote.serve_forever();break
            except OSError:
                REMOTE_ONLINE=False
                stop.wait(10)
    if config.get('enabled'):threading.Thread(target=vpn_listener,daemon=True).start()
    address='http://127.0.0.1:8765'
    print('JP Minato 1.36 · '+address,flush=True)
    if '--no-browser' not in sys.argv:webbrowser.open(address)
    try:server.serve_forever()
    except KeyboardInterrupt:pass
    finally:
        stop.set();REMOTE_ONLINE=False
        for remote in remote_servers:remote.shutdown();remote.server_close()
        for event in STORE.contact_cancel.values():event.set()
        STORE.inbox_stop.set();STORE.inbox_wake.set();STORE.whatsapp_stop.set();STORE.restoration_stop.set();STORE.restoration_wake.set()
        STORE.contact_engine.close();server.server_close()
        runtime.unlink(missing_ok=True)

if __name__=='__main__':main()
