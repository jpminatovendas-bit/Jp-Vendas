import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
import fitz
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import Store, norm

def workbook(city='Tijucas',phone='(48) 3263-0000'):
    w=openpyxl.Workbook();s=w.active
    s.append(['Código Único','Município','UF','Telefone Planilha Original','Telefones Coletados no Site','Presidente Atual (Nome)','Website Câmara'])
    s.append(['SC-TIJUCAS',city,'SC',phone,'(48) 3263-0000','Pessoa original','https://exemplo.sc.gov.br'])
    b=io.BytesIO();w.save(b);return b.getvalue()

def image(file):
    file.parent.mkdir(parents=True,exist_ok=True)
    doc=fitz.open();p=doc.new_page(width=160,height=200);p.draw_rect(p.rect,fill=(.2,.4,.6));p.get_pixmap().save(file);doc.close()

class Flow(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.store=Store(self.root/'data');self.store.import_xlsx(workbook());self.key='SC-TIJUCAS'
        self.originals=self.root/'originais'/'TIJUCAS'/'2004';self.restored=self.root/'tratadas'/'Tijucas'/'2004'
        self.store.set_setting('roots',{'SC':{'original':str(self.root/'originais'),'restored':str(self.root/'tratadas')}})
    def tearDown(self):self.temp.cleanup()
    def test_import_preserves_manual_edits(self):
        self.store.save_item(self.key,{'Presidente Atual (Nome)':'Nome corrigido'},{'sector':'Oeste','checked':True})
        result=self.store.import_xlsx(workbook(phone='(15) 1111-1111'))
        self.assertEqual(result['updated'],1)
        item=self.store.item(self.key)
        self.assertEqual(item['fields']['Presidente Atual (Nome)'],'Nome corrigido')
        self.assertEqual(item['local']['sector'],'Oeste');self.assertTrue(item['ddd_warning'])
        self.assertEqual(item['ddd'],'48')
    def test_folder_matching_and_pairing(self):
        image(self.originals/'JOÃO DA SILVA.jpg');image(self.restored/'JOAO_DA_SILVA.png')
        image(self.originals.parent/'2008'/'NOME DE OUTRO ANO.png')
        project=self.store.load_originals(self.key);self.assertEqual(len(project['pairs']),1)
        project=self.store.load_restored(self.key);self.assertIsNotNone(project['pairs'][0]['restored'])
        self.assertFalse(project['pairs'][0]['approved'])
    def test_duplicates_need_review(self):
        image(self.originals/'HOMONIMOS'/'1'/'MESMO NOME.png');image(self.originals/'HOMONIMOS'/'2'/'MESMO NOME.png')
        image(self.restored/'MESMO NOME.png')
        self.store.load_originals(self.key);project=self.store.load_restored(self.key)
        self.assertTrue(all(not x['restored'] for x in project['pairs']))
    def test_homonyms_match_subdirectories(self):
        for folder in ('1','2'):
            image(self.originals/'HOMONIMOS'/folder/'MESMO NOME.png');image(self.restored/'HOMONIMOS'/folder/'MESMO NOME.png')
        self.store.load_originals(self.key);p=self.store.load_restored(self.key)
        self.assertEqual(len({x['restored'] for x in p['pairs']}),2)
    def test_final_requires_restored_and_crest(self):
        image(self.originals/'PESSOA.png');self.store.load_originals(self.key)
        with self.assertRaises(ValueError):self.store.pdf(self.key)
        r=self.store.pdf(self.key,True);doc=fitz.open(self.store.files[r['url'].split('/')[-1]])
        self.assertIn('AGUARDANDO',doc[1].get_text());doc.close()
    def test_zip_preserves_names_and_selected_people(self):
        image(self.originals/'NOME SEM NUMERO.png');image(self.originals/'OUTRA PESSOA.png')
        p=self.store.load_originals(self.key)
        self.store.update_pairs(self.key,{'pairs':[{**x,'selected':x['name']=='NOME SEM NUMERO'} for x in p['pairs']]})
        r=self.store.export_originals(self.key)
        with zipfile.ZipFile(self.store.files[r['url'].split('/')[-1]]) as archive:
            self.assertEqual(set(archive.namelist()),{'NOME SEM NUMERO.png','manifest.json'})
    def test_final_allows_unapproved_complete_pairs(self):
        image(self.originals/'PESSOA.png');image(self.restored/'PESSOA.png')
        self.store.load_originals(self.key);p=self.store.load_restored(self.key)
        self.assertFalse(p['pairs'][0]['approved'])
        crest=self.root/'crest.png';image(crest);self.store.crest(self.key,crest.read_bytes(),'crest.png')
        result=self.store.pdf(self.key)
        self.assertEqual(result['pages'],9)
        self.store.update_pairs(self.key,{'pairs':[{'id':p['pairs'][0]['id'],'restored':p['pairs'][0]['restored'],'name':p['pairs'][0]['name'],'selected':True}]})
        self.assertFalse(self.store.project(self.key)['dirty'])
        self.store.save_item(self.key,{}, {'last_contact':'2026-10-10','commercial':'Contatado'})
        self.assertFalse(self.store.project(self.key)['dirty'])
        self.store.update_pairs(self.key,{'pairs':[{'id':p['pairs'][0]['id'],'restored':p['pairs'][0]['restored'],'name':'NOME ALTERADO','selected':True}]})
        self.assertTrue(self.store.project(self.key)['dirty'])
    def test_final_rejects_missing_restored_without_approval_gate(self):
        image(self.originals/'PESSOA.png');self.store.load_originals(self.key)
        crest=self.root/'crest.png';image(crest);self.store.crest(self.key,crest.read_bytes(),'crest.png')
        with self.assertRaisesRegex(ValueError,'tratadas'):self.store.pdf(self.key)
    def test_export_contacts_scope_and_literal_values(self):
        wb_other=openpyxl.load_workbook(io.BytesIO(workbook(city='Outra cidade')))
        wb_other.active.cell(2,1).value='SC-OUTRA';other=io.BytesIO();wb_other.save(other);wb_other.close();self.store.import_xlsx(other.getvalue())
        self.store.save_item(self.key,{'Presidente Atual (Nome)':'=EXTERNAL()'},{'sector':'Oeste','notes':'Notas comerciais'})
        result=self.store.export_contacts([self.key,'inexistente'])
        workbook_out=openpyxl.load_workbook(self.store.files[result['url'].split('/')[-1]])
        sheet=workbook_out.active
        self.assertEqual(sheet.max_row,2)
        headings=[c.value for c in sheet[1]]
        value=sheet.cell(2,headings.index('Presidente Atual (Nome)')+1)
        self.assertEqual(value.value,'=EXTERNAL()');self.assertEqual(value.data_type,'s')
        self.assertEqual(sheet.cell(2,headings.index('Setor comercial')+1).value,'Oeste')
        workbook_out.close()
        with self.assertRaises(ValueError):self.store.export_contacts([])
    def test_deleted_restored_blocks_final(self):
        image(self.originals/'PESSOA.png');image(self.restored/'PESSOA.png')
        self.store.load_originals(self.key);self.store.load_restored(self.key)
        crest=self.root/'crest.png';image(crest);self.store.crest(self.key,crest.read_bytes(),'crest.png')
        Path(self.store.project(self.key)['pairs'][0]['restored']).unlink()
        self.assertEqual(self.store.item(self.key)['stage'],'Em restauração')
        with self.assertRaisesRegex(ValueError,'movida ou apagada'):self.store.pdf(self.key)
    def test_browser_folder_upload_and_automatic_matching(self):
        source=self.root/'source.png';image(source);content=source.read_bytes()
        for kind,batch in [('original','a'*24),('restored','b'*24)]:
            self.store.photo_upload(self.key,kind,'2004/JOÃO DA SILVA.png',content,batch)
            project=self.store.finish_photo_upload(self.key,kind,batch)
        self.assertEqual(len(project['pairs']),1)
        self.assertIsNotNone(project['pairs'][0]['restored'])
        self.assertEqual(self.store.item(self.key)['stage'],'Pronto para PDF')
        self.store.photo_upload(self.key,'original','2004/JOÃO DA SILVA.png',content,'c'*24)
        again=self.store.finish_photo_upload(self.key,'original','c'*24)
        self.assertEqual(again['pairs'][0]['restored'],project['pairs'][0]['restored'])
    def test_browser_folder_upload_rejects_invalid_and_traversal(self):
        source=self.root/'source.png';image(source)
        for name in ['../escape.png','C:/escape.png','/escape.png']:
            with self.assertRaises(ValueError):self.store.photo_upload(self.key,'original',name,source.read_bytes(),'a'*24)
        with self.assertRaisesRegex(ValueError,'Imagem inválida'):self.store.photo_upload(self.key,'original','test.png',b'invalid','a'*24)
        self.assertEqual(self.store.project(self.key)['pairs'],[])
    def test_standard_upload_validation_and_order(self):
        image(self.originals/'PESSOA.png');image(self.restored/'PESSOA.png')
        self.store.load_originals(self.key);self.store.load_restored(self.key)
        crest=self.root/'crest.png';image(crest);self.store.crest(self.key,crest.read_bytes(),'crest.png')
        first=self.store.pdf(self.key);self.assertEqual(first['pages'],9)
        custom=fitz.open()
        for title in ['PADRAO PAGINA A','PADRAO PAGINA B']:
            page=custom.new_page();page.insert_text((20,40),title)
        info=self.store.upload_catalogue(custom.tobytes(),'novo.pdf');custom.close()
        self.assertEqual(info['pages'],2);self.assertTrue(self.store.project(self.key)['dirty'])
        result=self.store.pdf(self.key)
        with fitz.open(self.store.files[result['url'].split('/')[-1]]) as doc:
            self.assertEqual(len(doc),4)
            self.assertIn('TIJUCAS - SC',doc[0].get_text())
            self.assertNotIn('NOME DO MUNICIPIO',doc[0].get_text())
            self.assertIn('PESSOA',doc[1].get_text())
            self.assertIn('PADRAO PAGINA A',doc[2].get_text())
            self.assertIn('PADRAO PAGINA B',doc[3].get_text())
        with self.assertRaises(ValueError):self.store.upload_catalogue(b'not pdf','erro.pdf')
        self.assertEqual(self.store.catalogue_info()['filename'],'novo.pdf')
    def test_pdf_has_new_data_and_pagination(self):
        for i in range(10):
            image(self.originals/f'PESSOA TESTE {i}.png');image(self.restored/f'PESSOA TESTE {i}.png')
        self.store.load_originals(self.key);p=self.store.load_restored(self.key)
        self.assertTrue(all(not x['approved'] for x in p['pairs']))
        self.assertEqual(self.store.item(self.key)['stage'],'Pronto para PDF')
        crest=self.root/'crest.png';image(crest);self.store.crest(self.key,crest.read_bytes(),'crest.png')
        r=self.store.pdf(self.key);doc=fitz.open(self.store.files[r['url'].split('/')[-1]])
        self.assertEqual(len(doc),10);text=''.join(p.get_text() for p in list(doc)[1:3])
        self.assertIn('TIJUCAS',text);self.assertNotIn('FAXINALZINHO',text.upper());self.assertNotIn('AMÉRICO',text);self.assertNotIn('1989',text)
        self.assertTrue(all(f'PESSOA TESTE {i}' in text for i in range(10)))
        for page in doc:self.assertTrue(len(page.get_images())>0)
        doc.close();self.assertEqual(self.store.item(self.key)['stage'],'PDF pronto')

if __name__=='__main__':unittest.main()
